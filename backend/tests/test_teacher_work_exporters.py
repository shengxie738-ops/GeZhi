"""Detached T8 Office-byte tests. Four finite nodes; initially UNEXECUTED.

No DB, provider, application startup, browser, converter, host font, image decode,
chart, XLSX, extraction or fixture-file access. A separately reviewed immutable
runner must release actual Office imports/native execution. Missing reviewed
feature assertions are separate from import/environment errors.
"""
from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
from hashlib import sha256
from importlib import import_module
from io import BytesIO
import json
from pathlib import Path
import posixpath
from uuid import UUID
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo
from zlib import error as DeflateError

import pytest


BACKEND = Path(__file__).resolve().parents[1]
FEATURES = {
    "pptx": "app.services.teacher_work.exporters.pptx",
    "docx": "app.services.teacher_work.exporters.docx",
    "theme": "app.services.teacher_work.exporters.theme",
    "validation": "app.services.teacher_work.exporters.validation",
}
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
PPT_FONT = "字体：Microsoft YaHei；可替代：PingFang SC、黑体；未嵌入字体"
DOC_FONT = "字体：黑体/宋体；可替代：Microsoft YaHei、Songti SC；未嵌入字体"
TASK = UUID("10000000-0000-4000-8000-000000000001")
EVIDENCE = UUID("20000000-0000-4000-8000-000000000001")
REFERENCE = UUID("20000000-0000-4000-8000-000000000002")
NOW = datetime(2026, 10, 4, 0, 0, tzinfo=timezone.utc)
LAYOUTS = ("title", "section", "bullets", "two_column", "question", "summary")
FILE_CAP = 10 * 1024 * 1024


def _feature(kind):
    """Only absent approved product source is a clean feature assertion."""
    name = FEATURES[kind]
    path = BACKEND / (name.replace(".", "/") + ".py")
    assert path.is_file(), "MISSING_REVIEWED_FEATURE:" + name
    # Never catch ImportError: the runner must classify it as ENVIRONMENT_BLOCKED.
    return import_module(name)


def _api(first="pptx"):
    if first == "docx":
        doc = _feature("docx")
        ppt = _feature("pptx")
    else:
        ppt = _feature("pptx")
        doc = _feature("docx")
    theme = _feature("theme")
    validation = _feature("validation")
    return ppt.build_pptx, doc.build_docx, theme.ThemeVersion(), validation.validate_office_bytes


def _schema():
    return import_module("app.schemas.teacher_work")


def _slides(count=6):
    slides = []
    for index in range(count):
        layout = LAYOUTS[index % 6]
        slide = {
            "layout": layout, "title": f"第{index + 1}页：中文、标点与教学主题",
            "body": ["解释概念，不省略内容", "保留第一行\n保留第二行"], "columns": [],
            "notes": "" if index == 1 else f"讲者备注{index + 1}\n精确中文与引用",
            "source_note": "合成课件第3页；参考 DOI:10.0000/example（纯文本）",
            "evidence_refs": [EVIDENCE, REFERENCE],
        }
        if layout == "two_column":
            slide.update(body=[], columns=[["左列概念", "左列\n第二行"], ["右列比较"]])
        slides.append(slide)
    return _schema().SlideModel.model_validate({"slides": slides})


def _lesson(minutes=45):
    return _schema().LessonSnapshot.model_validate({
        "title": "合成中文教案", "topic": "中文教学主题", "course_name": "合成课程",
        "audience": "合成授课对象", "duration_minutes": minutes,
        "objectives": ["理解概念", "表达第一行\n表达第二行"],
        "key_points": ["教学重点"], "difficulties": ["教学难点"],
        "questions": ["课堂提问？"], "exercises": ["公开评价规则：解释依据与表达清晰"],
        "homework": ["课后任务"], "summary": "课件总结第一行\n课件总结第二行",
        "teaching_flow": [
            {"stage": "导入", "minutes": 10, "content": "导入第一行\n导入第二行"},
            {"stage": "讨论", "minutes": minutes - 10, "content": "讨论与练习"},
        ],
        "citations": [
            {"name": "合成课件", "page": 3, "excerpt": "真实冻结摘录\n第二行"},
            {"name": "参考 DOI:10.0000/example", "page": 0, "excerpt": "外部摘要"},
        ],
    })


def _version(model, lesson=None):
    schema = _schema()
    lesson = lesson or _lesson()
    sources = [
        {"evidence_id": EVIDENCE, "task_id": TASK, "resource_id": "synthetic-resource",
         "name": "合成课件", "page": 3, "excerpt": "真实冻结摘录",
         "resource_content_digest": "a" * 64, "acquired_at": NOW, "evidence_type": "courseware"},
        {"evidence_id": REFERENCE, "task_id": TASK,
         "ref_id": UUID("30000000-0000-4000-8000-000000000001"),
         "name": "参考 DOI:10.0000/example", "external_id": "10.0000/example",
         "excerpt": "外部摘要", "resource_content_digest": "b" * 64,
         "acquired_at": NOW, "evidence_type": "reference"},
    ]
    frozen = schema.FrozenPackageContent(lesson=lesson, slides=model.slides, source_snapshots=sources)
    payload = frozen.model_dump(mode="json")
    digest = sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True,
                               separators=(",", ":"), allow_nan=False).encode("utf-8")).hexdigest()
    return schema.PackageVersionDTO(
        **frozen.model_dump(), version_id=UUID("40000000-0000-4000-8000-000000000001"),
        task_id=TASK, version_no=1, run_id=UUID("50000000-0000-4000-8000-000000000001"),
        content_digest=digest, model_id="synthetic-model@1", skill_versions=["lesson_package@1"],
        exporter_versions=["gezhi-pptx@1", "gezhi-docx@1"],
        template_version="gezhi-office-theme@1", created_at=NOW,
    )


def _parts(raw):
    with ZipFile(BytesIO(raw)) as package:
        assert package.testzip() is None
        names = package.namelist()
        assert len(names) == len(set(names))
        return {name: package.read(name) for name in names}


def _xml(parts, name):
    return ET.fromstring(parts[name])


def _xml_bytes(root):
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _zip(parts, *, extra=(), compression=ZIP_DEFLATED):
    output = BytesIO()
    with ZipFile(output, "w", compression=compression) as package:
        seen = set()
        for name, content in list(parts.items()) + list(extra):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = compression
            if name in seen:
                with pytest.warns(UserWarning, match="Duplicate name"):
                    package.writestr(info, content)
            else:
                package.writestr(info, content)
            seen.add(name)
    return output.getvalue()


def _paragraph(element, dialect):
    text = []
    for child in element.iter():
        local = child.tag.rsplit("}", 1)[-1]
        if child.tag == "{" + NS[dialect] + "}t":
            text.append(child.text or "")
        elif local in ("br", "cr"):
            text.append("\n")
        elif local == "tab":
            text.append("\t")
    return "".join(text)


def _paragraphs(element, dialect):
    return [_paragraph(p, dialect) for p in element.findall(".//" + dialect + ":p", NS)]


def _relationship_target(owner, target):
    base = "" if owner == "" else posixpath.dirname(owner)
    return posixpath.normpath(posixpath.join(base, target)).lstrip("/")


def _relationships(parts, owner):
    name = "_rels/.rels" if not owner else posixpath.join(
        posixpath.dirname(owner), "_rels", posixpath.basename(owner) + ".rels")
    if name not in parts:
        return {}
    return {r.attrib["Id"]: r for r in _xml(parts, name)}


def _ordered_slides(parts):
    root = _xml(parts, "ppt/presentation.xml")
    relationships = _relationships(parts, "ppt/presentation.xml")
    return [_relationship_target("ppt/presentation.xml", relationships[s.attrib[
        "{" + NS["r"] + "}id"]].attrib["Target"])
        for s in root.findall("p:sldIdLst/p:sldId", NS)]


def _role_shapes(root):
    result = {}
    for shape in root.findall("p:cSld/p:spTree/p:sp", NS):
        props = shape.find("p:nvSpPr/p:cNvPr", NS)
        name = props.attrib.get("name", "")
        if name.startswith("gezhi:"):
            assert name not in result
            result[name] = shape
    return result


def _assert_graph(parts):
    """Independent OPC resolution; rIds are owner-local, never global."""
    content_types = _xml(parts, "[Content_Types].xml")
    defaults = {x.attrib["Extension"]: x.attrib["ContentType"] for x in content_types
                if x.tag.endswith("Default")}
    overrides = {x.attrib["PartName"].lstrip("/"): x.attrib["ContentType"]
                 for x in content_types if x.tag.endswith("Override")}
    for name in parts:
        assert name not in {"docProps/thumbnail.jpeg", "ppt/printerSettings/printerSettings1.bin"}
        assert not name.startswith(("customXml/", "ppt/media/", "word/media/", "ppt/embeddings/", "word/embeddings/"))
        assert name == "[Content_Types].xml" or name in overrides or name.rsplit(".", 1)[-1] in defaults
        if name.endswith(".rels"):
            if name == "_rels/.rels":
                owner = ""
            else:
                parent, filename = name.rsplit("/_rels/", 1)
                owner = parent + "/" + filename[:-5]
            ids = set()
            for edge in _xml(parts, name):
                assert edge.attrib["Id"] not in ids
                ids.add(edge.attrib["Id"])
                assert "TargetMode" not in edge.attrib
                assert _relationship_target(owner, edge.attrib["Target"]) in parts
                assert not any(term in edge.attrib["Type"].lower() for term in
                               ("hyperlink", "image", "thumbnail", "printersettings", "oleobject", "customxml"))
    assert "jpeg" not in defaults and "bin" not in defaults


def _assert_native_ppt(raw, model):
    parts = _parts(raw)
    _assert_graph(parts)
    presentation = _xml(parts, "ppt/presentation.xml")
    size = presentation.find("p:sldSz", NS)
    width, height = int(size.attrib["cx"]), int(size.attrib["cy"])
    assert width * 9 == height * 16
    ordered = _ordered_slides(parts)
    assert len(ordered) == len(model.slides)
    for number, (name, expected) in enumerate(zip(ordered, model.slides), 1):
        root = _xml(parts, name)
        assert not root.findall(".//p:pic", NS) and not root.findall(".//p:graphicFrame", NS)
        roles = _role_shapes(root)
        wanted = {"gezhi:title", "gezhi:source", "gezhi:font"}
        wanted |= {"gezhi:left", "gezhi:right"} if expected.layout == "two_column" else {"gezhi:body"}
        assert set(roles) == wanted
        assert _paragraphs(roles["gezhi:title"], "a") == [expected.title]
        assert _paragraphs(roles["gezhi:source"], "a") == [expected.source_note]
        assert _paragraphs(roles["gezhi:font"], "a") == [PPT_FONT]
        description = roles["gezhi:title"].find("p:nvSpPr/p:cNvPr", NS).attrib["descr"]
        assert json.loads(description) == {"layout": expected.layout,
            "evidence_refs": [str(item) for item in expected.evidence_refs]}
        bodies = [("gezhi:left", expected.columns[0]), ("gezhi:right", expected.columns[1])] \
            if expected.layout == "two_column" else [("gezhi:body", expected.body)]
        for role, items in bodies:
            assert _paragraphs(roles[role], "a") == list(items) or (not items and _paragraphs(roles[role], "a") == [""])
            body_pr = roles[role].find("p:txBody/a:bodyPr", NS)
            assert body_pr.find("a:normAutofit", NS) is None
            assert body_pr.find("a:spAutofit", NS) is None
            assert body_pr.find("a:noAutofit", NS) is not None
            for p in roles[role].findall("p:txBody/a:p", NS):
                props = p.findall("a:r/a:rPr", NS) + p.findall("a:pPr/a:defRPr", NS)
                assert props and all(int(pr.attrib["sz"]) >= 2000 for pr in props)
                assert all(pr.find("a:ea", NS).attrib["typeface"] == "Microsoft YaHei" for pr in props)
        for shape in root.findall("p:cSld/p:spTree/p:sp", NS):
            transform = shape.find("p:spPr/a:xfrm", NS)
            assert transform is not None
            offset, extent = transform.find("a:off", NS), transform.find("a:ext", NS)
            x, y, cx, cy = [int(element.attrib[key]) for element, key in
                            ((offset, "x"), (offset, "y"), (extent, "cx"), (extent, "cy"))]
            assert x >= 0 and y >= 0 and cx > 0 and cy > 0 and x + cx <= width and y + cy <= height
        note_edges = [r for r in _relationships(parts, name).values() if r.attrib["Type"].endswith("/notesSlide")]
        assert len(note_edges) == 1
        notes = _xml(parts, _relationship_target(name, note_edges[0].attrib["Target"]))
        body_shapes = [s for s in notes.findall("p:cSld/p:spTree/p:sp", NS)
                       if s.find("p:nvSpPr/p:nvPr/p:ph", NS) is not None and
                       s.find("p:nvSpPr/p:nvPr/p:ph", NS).attrib.get("type") == "body"]
        assert len(body_shapes) == 1
        assert "\n".join(_paragraphs(body_shapes[0], "a")) == expected.notes, number
    # Reopen only generated safe bytes; reading notes_slide never fabricates absent notes.
    from pptx import Presentation
    reopened = Presentation(BytesIO(raw))
    assert len(reopened.slides) == len(model.slides)
    for slide, expected in zip(reopened.slides, model.slides):
        assert slide.has_notes_slide
        assert slide.notes_slide.notes_text_frame.text.replace("\v", "\n") == expected.notes
        assert [s.text_frame.paragraphs[0].text for s in slide.shapes if s.name == "gezhi:title"] == [expected.title]


def _doc_projection(lesson):
    values = [lesson.title, (lesson.course_name + " 教案") if lesson.course_name else "教案",
        "课程名称", lesson.course_name or "—", "授课对象", lesson.audience or "—",
        "备课主题", lesson.topic or "—", "课时", f"{lesson.duration_minutes} 分钟"]
    sections = []
    for field, heading in (("objectives", "教学目标"), ("key_points", "教学重点"), ("difficulties", "教学难点")):
        if getattr(lesson, field):
            sections.append((heading, [f"{i}. {t}" for i, t in enumerate(getattr(lesson, field), 1)]))
    flow = ["教学环节", "时长(分钟)", "教学内容"]
    for stage in lesson.teaching_flow:
        flow.extend([stage.stage, str(stage.minutes), stage.content])
    sections.append(("教学流程", flow))
    for field, heading in (("questions", "课堂提问"), ("exercises", "课堂练习"), ("homework", "课后作业")):
        if getattr(lesson, field):
            sections.append((heading, [f"{i}. {t}" for i, t in enumerate(getattr(lesson, field), 1)]))
    if lesson.summary:
        sections.append(("课件总结", lesson.summary.splitlines()))
    if lesson.citations:
        citations = []
        for i, citation in enumerate(lesson.citations, 1):
            citations.append(f"{i}. {citation.name or '未命名课件'}" + (f" 第{citation.page}页" if citation.page else ""))
            if citation.excerpt:
                citations.append("摘录：" + citation.excerpt)
        sections.append(("课件依据", citations))
    for i, (heading, content) in enumerate(sections):
        values.append("一二三四五六七八九十"[i] + "、" + heading)
        values.extend(content)
    return values + [DOC_FONT]


def _assert_doc(raw, lesson):
    parts = _parts(raw)
    _assert_graph(parts)
    root = _xml(parts, "word/document.xml")
    assert _paragraphs(root, "w") == _doc_projection(lesson)
    fonts = root.findall(".//w:rPr/w:rFonts", NS)
    assert fonts and {f.attrib["{" + NS["w"] + "}eastAsia"] for f in fonts} >= {"黑体", "宋体"}
    from docx import Document
    reopened = Document(BytesIO(raw))
    texts = []
    for child in reopened.element.body:
        for paragraph in child.iter("{" + NS["w"] + "}p"):
            texts.append(_paragraph(paragraph, "w"))
    assert texts == _doc_projection(lesson)
    assert [cell.text for row in reopened.tables[1].rows[1:] for cell in row.cells] == [
        text for stage in lesson.teaching_flow for text in (stage.stage, str(stage.minutes), stage.content)]
    assert sum(int(row.cells[1].text) for row in reopened.tables[1].rows[1:]) == lesson.duration_minutes


def _valid(validator, kind, raw, version):
    result = validator(kind, raw, version)
    assert isinstance(result, _schema().ValidationSummary)
    assert result.valid, (kind, result)
    assert result.checks and any("实际" in warning or "Office" in warning for warning in result.warnings)
    return result


def test_native_editable_chinese_pptx():
    """Deleting native text/notes/order/font/digest fidelity must fail this node."""
    build_pptx, _, theme, validate = _api()
    for count in (6, 8, 12):
        model = _slides(count)
        before = model.model_dump(mode="json")
        version = _version(model)
        version_before = version.model_dump(mode="json")
        raw = build_pptx(model, theme)
        assert type(raw) is bytes and 0 < len(raw) <= FILE_CAP
        _assert_native_ppt(raw, model)
        _valid(validate, "pptx", raw, version)
        assert model.model_dump(mode="json") == before
        assert version.model_dump(mode="json") == version_before
        if count == 8:
            assert build_pptx(model, theme) == raw  # Fixed ZIP/core metadata, not semantic-only determinism.
    with pytest.raises((AttributeError, TypeError)):
        theme.version = "untrusted-template"


def test_slide_density_fails_without_clipping():
    """Clipping, dropping, adding pages or shrinking dense text must fail."""
    build_pptx, _, theme, _ = _api()
    base = _slides().model_dump()
    variants = (
        {"title": "题\n" * 29},
        {"body": ["行\n" * 44]},
        {"body": ["密" * 90] * 4},
        {"layout": "two_column", "body": [], "columns": [["密" * 90] * 2, ["密" * 90] * 2]},
    )
    for change in variants:
        payload = deepcopy(base)
        payload["slides"][1].update(change)
        model = _schema().SlideModel.model_validate(payload)
        before = model.model_dump(mode="json")
        with pytest.raises(ValueError, match=r"SLIDE_OVERFLOW.*(?:page=2|:2|页2)"):
            build_pptx(model, theme)
        assert model.model_dump(mode="json") == before
    # Bounded Chinese and meaningful line breaks remain editable, without fitting APIs.
    _assert_native_ppt(build_pptx(_slides(), theme), _slides())


def test_docx_frozen_mapping():
    """Using a live lesson, losing tables/whitespace or accepting bypassed bounds fails."""
    _, build_docx, _, validate = _api("docx")
    for minutes in (45, 60):
        lesson = _lesson(minutes)
        model = _slides()
        version = _version(model, lesson)
        before = lesson.model_dump(mode="json")
        raw = build_docx(lesson)
        _assert_doc(raw, lesson)
        _valid(validate, "docx", raw, version)
        assert lesson.model_dump(mode="json") == before
        assert build_docx(lesson) == raw
    # Strict Work rejects cases the old permissive body would strip/drop or cannot encode.
    base = _lesson().model_dump()
    changes = (
        {"title": "  meaningful boundary  "}, {"objectives": ["   "]},
        {"summary": "第一行\n\n第二行"}, {"summary": "第一行\n 第二行"},
        {"citations": [{"name": "", "page": 0, "excerpt": ""}]},
        {"questions": ["不可编码\x01内容"]},
    )
    for change in changes:
        lesson = _schema().LessonSnapshot.model_validate({**deepcopy(base), **change})
        with pytest.raises(ValueError, match="DOCX_UNREPRESENTABLE"):
            build_docx(lesson)
    for change in ({"duration_minutes": 46}, {"objectives": ["越" * 2001]}):
        bypassed = _lesson().model_copy(update=change)
        with pytest.raises(ValueError):
            build_docx(bypassed)


def _mutate_xml(parts, name, transform):
    copied = dict(parts)
    root = _xml(parts, name)
    transform(root)
    copied[name] = _xml_bytes(root)
    return _zip(copied)


def _bad_edge(parts, owner, **attributes):
    name = "_rels/.rels" if not owner else posixpath.join(posixpath.dirname(owner), "_rels", posixpath.basename(owner) + ".rels")
    root = _xml(parts, name)
    edge = root[0]
    edge.attrib.update(attributes)
    copied = dict(parts)
    copied[name] = _xml_bytes(root)
    return _zip(copied)


def _parser_regressions(parts):
    """Three bounded, source-supported malformed ZIP/DEFLATE mutations."""
    raw = _zip(parts)
    with ZipFile(BytesIO(raw)) as package:
        member = package.infolist()[0]
        assert member.header_offset == 0 and member.compress_type == ZIP_DEFLATED and member.compress_size > 0
    assert raw[:4] == b"PK\x03\x04" and raw[-22:-18] == b"PK\x05\x06"
    central = int.from_bytes(raw[-6:-2], "little")
    assert raw[central:central + 4] == b"PK\x01\x02"
    for bit in (5, 6):
        changed = bytearray(raw)
        for offset in (6, central + 8):
            flags = int.from_bytes(changed[offset:offset + 2], "little") | (1 << bit)
            changed[offset:offset + 2] = flags.to_bytes(2, "little")
        yield "zip_flag_bit" + str(bit), bytes(changed)
    payload = 30 + int.from_bytes(raw[26:28], "little") + int.from_bytes(raw[28:30], "little")
    changed = bytearray(raw)
    # Raw DEFLATE BFINAL=1/BTYPE=3 is reserved. The rest of the ZIP stays intact.
    changed[payload] = 0x07
    yield "deflate_reserved_block", bytes(changed)


def _group_extra_text(root):
    """A real native grouped textbox ignored by a direct-sp-only projection."""
    group = ET.Element("{" + NS["p"] + "}grpSp")
    nonvisual = ET.SubElement(group, "{" + NS["p"] + "}nvGrpSpPr")
    ET.SubElement(nonvisual, "{" + NS["p"] + "}cNvPr", {"id": "99", "name": "unexpected-group"})
    ET.SubElement(nonvisual, "{" + NS["p"] + "}cNvGrpSpPr")
    ET.SubElement(nonvisual, "{" + NS["p"] + "}nvPr")
    props = ET.SubElement(group, "{" + NS["p"] + "}grpSpPr")
    transform = ET.SubElement(props, "{" + NS["a"] + "}xfrm")
    for tag, attributes in (("off", {"x": "0", "y": "0"}),
                            ("ext", {"cx": "12192000", "cy": "6858000"}),
                            ("chOff", {"x": "0", "y": "0"}),
                            ("chExt", {"cx": "12192000", "cy": "6858000"})):
        ET.SubElement(transform, "{" + NS["a"] + "}" + tag, attributes)
    child = deepcopy(_role_shapes(root)["gezhi:body"])
    child.find("p:nvSpPr/p:cNvPr", NS).attrib.update({"id": "100", "name": "unexpected-group-text"})
    child.find(".//a:t", NS).text = "冻结模型之外的额外原生文字"
    group.append(child)
    root.find("p:cSld/p:spTree", NS).append(group)


def _hostile(parts, kind):
    """Exact finite synthetic ZIP/XML mutation inventory; no external execution."""
    main = "ppt/presentation.xml" if kind == "pptx" else "word/document.xml"
    owner_rels = "ppt/_rels/presentation.xml.rels" if kind == "pptx" else "word/_rels/document.xml.rels"
    yield from _parser_regressions(parts)
    yield "not_zip", b"not an office package"
    yield "empty", b""
    yield "over_file_cap", b"x" * (FILE_CAP + 1)
    yield "missing_main", _zip({n: b for n, b in parts.items() if n != main})
    yield "dangling", _bad_edge(parts, main, Target="missing.xml")
    yield "external", _bad_edge(parts, main, Target="https://example.invalid/payload", TargetMode="External")
    yield "scheme_internal", _bad_edge(parts, main, Target="file:///private/payload")
    yield "root_escape", _bad_edge(parts, main, Target="../../../../outside.xml")
    yield "hyperlink_type", _bad_edge(parts, main, Type=NS["r"] + "/hyperlink")
    yield "duplicate_relationship", _mutate_xml(parts, owner_rels, lambda root: root.append(deepcopy(root[0])))
    yield "bad_content_type", _mutate_xml(parts, "[Content_Types].xml", lambda root: root[0].set("ContentType", "application/x-executable"))
    for name, data in (
        ("../escape.xml", b"<x/>"), ("/absolute.xml", b"<x/>"), ("C:/drive.xml", b"<x/>"),
        ("ppt/vbaProject.bin", b"macro"), ("word/embeddings/oleObject1.bin", b"object"),
        ("word/media/image1.png", b"image"), ("customXml/unexpected.xml", b"<x/>"),
        ("docProps/thumbnail.jpeg", b"changed vendor thumbnail"),
        ("ppt/printerSettings/printerSettings1.bin", b"changed vendor printer settings"),
        ("unknown.xml", b"<x/>"),
    ):
        yield name, _zip(parts, extra=((name, data),))
    # Duplicate archive members and noncanonical aliases are independently rejected.
    yield "duplicate_member", _zip(parts, extra=((main, parts[main]),))
    yield "alias_member", _zip(parts, extra=((main.replace("/", "/./", 1), parts[main]),))
    yield "oversized_xml", _zip(parts, extra=(("large.xml", b"<x>" + b"a" * (2 * 1024 * 1024) + b"</x>"),))
    yield "member_cap", _zip(parts, extra=tuple((f"extra/{i}.xml", b"<x/>") for i in range(257)))
    for encoding in ("utf-8", "utf-16", "utf-32"):
        hostile = ('<?xml version="1.0" encoding="' + encoding + '"?>'
                   '<!DOCTYPE x [<!ENTITY payload SYSTEM "file:///private/value">]><x>&payload;</x>')
        copied = dict(parts)
        copied[main] = hostile.encode(encoding)
        yield "dtd_" + encoding, _zip(copied)
    copied = dict(parts)
    copied[main] = b"<broken"
    yield "malformed_xml", _zip(copied)
    if kind == "pptx":
        slide = _ordered_slides(parts)[0]
        note = next(r for r in _relationships(parts, slide).values() if r.attrib["Type"].endswith("/notesSlide"))
        note_name = _relationship_target(slide, note.attrib["Target"])
        yield "missing_notes", _zip({n: b for n, b in parts.items() if n != note_name})
        yield "extra_slide_text", _mutate_xml(parts, slide, lambda root: root.find(".//a:t", NS).__setattr__("text", "changed"))
        yield "body_below_floor", _mutate_xml(parts, slide, lambda root: _role_shapes(root)["gezhi:body"].find("p:txBody/a:p/a:r/a:rPr", NS).set("sz", "1900"))
        yield "off_canvas", _mutate_xml(parts, slide, lambda root: _role_shapes(root)["gezhi:body"].find("p:spPr/a:xfrm/a:off", NS).set("x", "999999999"))
        yield "shrink_to_fit", _mutate_xml(parts, slide, lambda root: _role_shapes(root)["gezhi:body"].find("p:txBody/a:bodyPr", NS).append(ET.Element("{" + NS["a"] + "}normAutofit")))
        yield "animation", _mutate_xml(parts, slide, lambda root: root.append(ET.Element("{" + NS["p"] + "}timing")))
        yield "extra_grouped_text", _mutate_xml(parts, slide, _group_extra_text)
    else:
        yield "changed_lesson", _mutate_xml(parts, main, lambda root: root.find(".//w:t", NS).__setattr__("text", "changed"))
        yield "external_field", _mutate_xml(parts, main, lambda root: ET.SubElement(root.find("w:body/w:p", NS), "{" + NS["w"] + "}fldSimple", {"{" + NS["w"] + "}instr": 'INCLUDETEXT "https://example.invalid/content"'}))
        yield "alt_chunk", _mutate_xml(parts, main, lambda root: root.find("w:body", NS).append(ET.Element("{" + NS["w"] + "}altChunk")))


def test_ooxml_rejects_hostile_parts():
    """Accepting any unsafe/changed package or laundering changed bootstrap fails."""
    build_pptx, build_docx, theme, validate = _api()
    model, lesson = _slides(), _lesson()
    version = _version(model, lesson)
    regressions = []
    parser_labels = {"zip_flag_bit5", "zip_flag_bit6", "deflate_reserved_block"}
    for kind, raw in (("pptx", build_pptx(model, theme)), ("docx", build_docx(lesson))):
        _valid(validate, kind, raw, version)
        parts = _parts(raw)
        for label, hostile in _hostile(parts, kind):
            if label in parser_labels:
                try:
                    result = validate(kind, hostile, version)
                except (NotImplementedError, EOFError, DeflateError) as error:
                    exception_name = "zlib.error" if isinstance(error, DeflateError) else type(error).__name__
                    regressions.append("PARSER_ESCAPE:" + kind + ":" + label + ":" + exception_name)
                    continue
            else:
                # Never catch import, policy or unrelated runtime/environment failures.
                result = validate(kind, hostile, version)
            assert isinstance(result, _schema().ValidationSummary)
            if label == "extra_grouped_text" and result.valid:
                regressions.append("ACCEPTED_GROUPED_EXTRA_TEXT:" + kind)
                continue
            assert not result.valid, (kind, label)
            assert result.checks or result.warnings, (kind, label)
        for changed in (
            version.model_copy(update={"content_digest": "0" * 64}),
            version.model_copy(update={"slides": (model.slides[0].model_copy(update={"evidence_refs": (UUID("99999999-0000-4000-8000-000000000001"),)}), *model.slides[1:])}),
            version.model_copy(update={"source_snapshots": (version.source_snapshots[0].model_copy(update={"task_id": UUID("99999999-0000-4000-8000-000000000002")}), version.source_snapshots[1])}),
        ):
            assert not validate(kind, raw, changed).valid
        assert not validate("pdf", raw, version).valid
        assert not validate(kind, bytearray(raw), version).valid
    # Exact trusted factories only: no hostile package reaches a convenience reader.
    finalizer = _feature("validation")._finalize_generated_office_bytes
    from pptx import Presentation
    from docx import Document
    for kind, factory in (("pptx", Presentation), ("docx", Document)):
        buffer = BytesIO()
        factory().save(buffer)
        trusted = buffer.getvalue()
        cleaned = finalizer(kind, trusted)
        _assert_graph(_parts(cleaned))
        assert finalizer(kind, cleaned) == cleaned  # Explicit already-absent idempotence.
        pinned = _parts(trusted)
        changed = dict(pinned)
        changed["docProps/thumbnail.jpeg"] += b"changed"
        with pytest.raises(ValueError):
            finalizer(kind, _zip(changed))
        with pytest.raises(ValueError):
            finalizer(kind, _mutate_xml(pinned, "_rels/.rels", lambda root: next(
                edge for edge in root if edge.attrib["Type"].endswith("/metadata/thumbnail")).set("Id", "changed-owner-id")))
        if kind == "pptx":
            changed = dict(pinned)
            changed["ppt/printerSettings/printerSettings1.bin"] += b"changed"
            with pytest.raises(ValueError):
                finalizer(kind, _zip(changed))
        else:
            with pytest.raises(ValueError):
                finalizer(kind, _zip({n: b for n, b in pinned.items() if n != "customXml/itemProps1.xml"}))
            comment = _mutate_xml(pinned, "customXml/_rels/item1.xml.rels",
                                  lambda root: root.append(ET.Comment("unexpected bibliography node")))
            try:
                finalizer(kind, comment)
            except ValueError:
                pass
            else:
                regressions.append("LAUNDERED_BIBLIOGRAPHY_COMMENT:" + kind)
    assert not regressions, "OOXML_REGRESSION_FAILURES:" + ",".join(regressions)
