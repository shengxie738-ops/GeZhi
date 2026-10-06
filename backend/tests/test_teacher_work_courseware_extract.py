"""Narrow, offline parser tests; no app, provider, database, or service startup.

Run with: python -m unittest discover -s backend/tests \
    -p test_teacher_work_courseware_extract.py -v
"""
from __future__ import annotations

from dataclasses import FrozenInstanceError, replace
from hashlib import sha256
from importlib import import_module
from io import BytesIO
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
try:
    subject = import_module("app.services.teacher_work.courseware_extract")
except ModuleNotFoundError as error:
    if error.name != "app.services.teacher_work.courseware_extract":
        raise
    subject = None


def pdf_bytes(texts, *, encrypted=False):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream.set_data(f"BT /F1 12 Tf 30 740 Td ({escaped}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("secret")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def pptx_bytes(texts, *, table=False, group=False):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    for text in texts:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        if text:
            shapes = slide.shapes.add_group_shape().shapes if group else slide.shapes
            shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = text
        if table:
            cells = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
            cells.cell(0, 0).text = "alpha"
            cells.cell(0, 1).text = "beta"
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def zip_bytes(entries):
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, value in entries:
            archive.writestr(name, value)
    return output.getvalue()


class FakeTrustedRunner:
    """Contract double only: never proof that this test process is isolated."""
    def __init__(self, pages=(), *, support=True, error=None):
        self.pages, self.support, self.error = pages, support, error
        self.calls = []

    def verify_support(self, limits):
        return self.support

    def run(self, data, extension, limits):
        self.calls.append((data, extension, limits))
        if self.error:
            raise self.error
        return self.pages


class CoursewareExtractionTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(subject, "bounded courseware extraction module is not implemented")

    def assert_code(self, code, data, extension, limits=None):
        with self.assertRaises(subject.ExtractionError) as caught:
            subject.extract_pages_untrusted(data, extension, limits or subject.DEFAULT_LIMITS)
        self.assertEqual(caught.exception.code, code)
        self.assertEqual(str(caught.exception), code)

    def test_frozen_contract_and_conservative_limit_validation(self):
        with self.assertRaises(FrozenInstanceError):
            subject.DEFAULT_LIMITS.max_pages = 500
        with self.assertRaises(FrozenInstanceError):
            subject.ExtractedPage(1, "text").page = 2
        for updates in ({"max_pages": 0}, {"max_pages": True}, {"max_pages": 201},
                        {"max_wall_seconds": float("inf")}, {"max_file_bytes": "10"}):
            with self.subTest(updates=updates), self.assertRaises(ValueError):
                replace(subject.DEFAULT_LIMITS, **updates)

    def test_pdf_preserves_physical_blank_page_numbers(self):
        pages = subject.extract_pages_untrusted(pdf_bytes(["first", "", "third"]), ".pdf")
        self.assertEqual([p.page for p in pages], [1, 2, 3])
        self.assertEqual([p.text.strip() for p in pages], ["first", "", "third"])
        self.assertIs(type(pages), tuple)

    def test_pptx_preserves_slide_numbers_and_unicode(self):
        pages = subject.extract_pages_untrusted(pptx_bytes(["第一张", "", "Third"]), ".PPTX")
        self.assertEqual([(p.page, p.text.strip()) for p in pages], [(1, "第一张"), (2, ""), (3, "Third")])

    def test_pptx_extracts_group_text_and_table_cells(self):
        pages = subject.extract_pages_untrusted(pptx_bytes(["group"], table=True, group=True), ".pptx")
        for expected in ["group", "alpha", "beta"]:
            self.assertIn(expected, pages[0].text)

    def test_rejects_empty_image_only_and_encrypted_sources(self):
        self.assert_code("EMPTY_INPUT", b"", ".pdf")
        self.assert_code("NO_EXTRACTABLE_TEXT", pdf_bytes([""]), ".pdf")
        self.assert_code("NO_EXTRACTABLE_TEXT", pptx_bytes([""]), ".pptx")
        self.assert_code("ENCRYPTED_SOURCE", pdf_bytes(["private"], encrypted=True), ".pdf")

    def test_rejects_paths_mutable_bytes_legacy_ppt_and_wrong_magic(self):
        for data in ["/private/source.pdf", BytesIO(b"%PDF"), bytearray(b"%PDF")]:
            self.assert_code("INVALID_INPUT", data, ".pdf")
        for extension in [".ppt", "pdf", "source.pdf", ".docx", None]:
            self.assert_code("UNSUPPORTED_FORMAT", b"%PDF-1.7\n", extension)
        self.assert_code("FORMAT_MISMATCH", b"not a pdf", ".pdf")
        self.assert_code("FORMAT_MISMATCH", pdf_bytes(["text"]), ".pptx")

    def test_corrupt_sources_fail_with_sanitized_codes(self):
        self.assert_code("CORRUPT_SOURCE", b"%PDF-1.7\nprivate-error-data\n%%EOF", ".pdf")
        self.assert_code("CORRUPT_SOURCE", b"PK\x03\x04private-error-data", ".pptx")
        self.assert_code("CORRUPT_SOURCE", zip_bytes([("a.xml", "<a/>")]), ".pptx")

    def test_file_page_and_text_limits_fail_without_partial_success(self):
        data = pdf_bytes(["first", "second"])
        self.assert_code("FILE_LIMIT", data, ".pdf", replace(subject.DEFAULT_LIMITS, max_file_bytes=len(data)-1))
        for data, extension in [(data, ".pdf"), (pptx_bytes(["first", "second"]), ".pptx")]:
            self.assert_code("PAGE_LIMIT", data, extension, replace(subject.DEFAULT_LIMITS, max_pages=1))
            self.assert_code("PAGE_TEXT_LIMIT", data, extension, replace(subject.DEFAULT_LIMITS, max_page_chars=3))
            self.assert_code("TOTAL_TEXT_LIMIT", data, extension, replace(subject.DEFAULT_LIMITS, max_total_chars=7))

    def test_whitespace_aggregate_budget_stops_before_later_page_parse(self):
        class Page:
            def __init__(self, text):
                self.text = text
            def extract_text(self):
                if self.text is None:
                    raise AssertionError("a later page must not be parsed after budget overflow")
                return self.text
        class Reader:
            is_encrypted = False
            pages = [Page("    "), Page("    "), Page(None)]
        for field, code in [("max_total_chars", "TOTAL_TEXT_LIMIT"), ("max_output_bytes", "OUTPUT_LIMIT")]:
            with self.subTest(field=field), patch("pypdf.PdfReader", return_value=Reader()):
                self.assert_code(code, b"%PDF-1.7\n%%EOF", ".pdf", replace(subject.DEFAULT_LIMITS, **{field: 7}))

    def test_output_byte_budget_counts_utf8(self):
        self.assert_code("OUTPUT_LIMIT", pptx_bytes(["汉字"]), ".pptx",
                         replace(subject.DEFAULT_LIMITS, max_output_bytes=5))

    def test_zip_member_count_size_total_and_compression_budgets(self):
        self.assert_code("ZIP_ENTRY_LIMIT", zip_bytes([("a", "a"), ("b", "b")]), ".pptx",
                         replace(subject.DEFAULT_LIMITS, max_zip_entries=1))
        self.assert_code("ZIP_MEMBER_LIMIT", zip_bytes([("a", "a"*11)]), ".pptx",
                         replace(subject.DEFAULT_LIMITS, max_zip_member_bytes=10))
        self.assert_code("ZIP_TOTAL_LIMIT", zip_bytes([("a", "abcdef"), ("b", "ghijkl")]), ".pptx",
                         replace(subject.DEFAULT_LIMITS, max_zip_total_bytes=10))
        self.assert_code("ZIP_RATIO_LIMIT", zip_bytes([("a", "a"*10000)]), ".pptx")

    def test_zip_unsafe_names_and_duplicates_are_rejected(self):
        for name in ["../escape", "/absolute", "a/../b", "a\\b"]:
            with self.subTest(name=name):
                self.assert_code("UNSAFE_ARCHIVE", zip_bytes([(name, "data")]), ".pptx")
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            duplicate = zip_bytes([("a.xml", "<a/>"), ("a.xml", "<a/>")])
        self.assert_code("UNSAFE_ARCHIVE", duplicate, ".pptx")

    def test_pptx_dtd_entities_and_deep_xml_are_rejected(self):
        self.assert_code("UNSAFE_XML", zip_bytes([("a.xml", '<!DOCTYPE a [<!ENTITY x "secret">]><a>&x;</a>')]), ".pptx")
        self.assert_code("XML_DEPTH_LIMIT", zip_bytes([("a.xml", "<a>"*8+"</a>"*8)]), ".pptx",
                         replace(subject.DEFAULT_LIMITS, max_xml_depth=4))
        self.assert_code("XML_ELEMENT_LIMIT", zip_bytes([("a.xml", "<a><b/><c/></a>")]), ".pptx",
                         replace(subject.DEFAULT_LIMITS, max_xml_elements=2))

    def test_xml_budgets_follow_content_types_for_disguised_parts(self):
        manifest = ('<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                    '<Override PartName="/disguised.bin" ContentType="application/custom+xml"/>'
                    '</Types>')
        disguised = '<!DOCTYPE a [<!ENTITY x "hidden">]><a>&x;</a>'
        self.assert_code("UNSAFE_XML", zip_bytes([("[Content_Types].xml", manifest),
                                                 ("disguised.bin", disguised)]), ".pptx")

    def test_xml_preflight_rejects_utf16_entity_declarations(self):
        source = '<?xml version="1.0" encoding="utf-16"?><!DOCTYPE a [<!ENTITY x "hidden">]><a>&x;</a>'
        self.assert_code("UNSAFE_XML", zip_bytes([("part.XML", source.encode("utf-16"))]), ".pptx")

    def test_default_supervision_fails_closed_without_running_parser(self):
        data = pdf_bytes(["text"])
        with patch.object(subject, "extract_pages_untrusted", side_effect=AssertionError("must not parse")):
            result = subject.extract_pages(data, ".pdf")
        self.assertEqual((result.status, result.error_code, result.pages), ("unavailable", "ISOLATION_UNAVAILABLE", ()))
        self.assertEqual(result.content_sha256, sha256(data).hexdigest())

    def test_runner_requires_verified_support_and_exact_true(self):
        for support in [False, None, 1, "yes"]:
            with self.subTest(support=support):
                runner = FakeTrustedRunner(support=support)
                result = subject.extract_pages(pdf_bytes(["text"]), ".pdf", runner=runner)
                self.assertEqual(result.status, "unavailable")
                self.assertEqual(runner.calls, [])

    def test_supervised_result_hashes_exact_bytes_and_forwards_fixed_limits(self):
        data = pdf_bytes(["text"])
        pages = (subject.ExtractedPage(1, "text"),)
        runner = FakeTrustedRunner(pages)
        result = subject.extract_pages(data, ".PDF", runner=runner)
        self.assertEqual(result.status, "ready")
        self.assertIsNone(result.error_code)
        self.assertEqual(result.pages, pages)
        self.assertEqual(result.content_sha256, sha256(data).hexdigest())
        self.assertEqual(runner.calls, [(data, ".pdf", subject.DEFAULT_LIMITS)])

    def test_runner_output_revalidated_as_bounded_contiguous_immutable_pages(self):
        invalid = [[], (), [subject.ExtractedPage(1, "a")],
                   (subject.ExtractedPage(0, "a"),), (subject.ExtractedPage(True, "a"),),
                   (subject.ExtractedPage(2, "a"),),
                   (subject.ExtractedPage(1, "a"), subject.ExtractedPage(1, "b")),
                   (subject.ExtractedPage(1, ""),), (subject.ExtractedPage(1, None),),
                   (subject.ExtractedPage(1, "a"*20001),)]
        for pages in invalid:
            with self.subTest(pages_type=type(pages), count=len(pages)):
                result = subject.extract_pages(pdf_bytes(["text"]), ".pdf", runner=FakeTrustedRunner(pages))
                self.assertEqual(result.status, "failed")
                self.assertEqual(result.pages, ())

    def test_exact_page_missing_fields_is_sanitized(self):
        empty = object.__new__(subject.ExtractedPage)
        missing_text = object.__new__(subject.ExtractedPage)
        object.__setattr__(missing_text, "page", 1)
        missing_page = object.__new__(subject.ExtractedPage)
        object.__setattr__(missing_page, "text", "synthetic")
        for page in (empty, missing_text, missing_page):
            with self.subTest(fields=tuple(vars(page))):
                result = subject.extract_pages(pdf_bytes(["synthetic"]), ".pdf",
                                               runner=FakeTrustedRunner((page,)))
                self.assertEqual((result.status, result.error_code, result.pages),
                                 ("failed", "INVALID_OUTPUT", ()))

    def test_mutated_missing_and_nonhashable_error_codes_are_sanitized(self):
        class CodeSubclass(str):
            pass
        for value in ("synthetic-private-detail", None, [], {}, 1, CodeSubclass("CORRUPT_SOURCE")):
            with self.subTest(value_type=type(value)):
                error = subject.ExtractionError("CORRUPT_SOURCE")
                error.code = value
                result = subject.extract_pages(pdf_bytes(["synthetic"]), ".pdf",
                                               runner=FakeTrustedRunner(error=error))
                self.assertEqual((result.status, result.error_code, result.pages),
                                 ("failed", "EXTRACTION_FAILED", ()))
        error = subject.ExtractionError("CORRUPT_SOURCE")
        del error.code
        result = subject.extract_pages(pdf_bytes(["synthetic"]), ".pdf",
                                       runner=FakeTrustedRunner(error=error))
        self.assertEqual((result.status, result.error_code, result.pages),
                         ("failed", "EXTRACTION_FAILED", ()))

    def test_failed_input_does_not_launch_runner(self):
        runner = FakeTrustedRunner()
        result = subject.extract_pages(b"not PDF", ".pdf", runner=runner)
        self.assertEqual((result.status, result.error_code), ("failed", "FORMAT_MISMATCH"))
        self.assertEqual(runner.calls, [])
        self.assertEqual(subject.extract_pages("private/path.pdf", ".pdf").content_sha256, "")

    def test_runner_failure_is_sanitized_without_fallback(self):
        for error, code in [(RuntimeError("private path or credentials"), "RUNNER_FAILED"),
                            (TimeoutError("worker still alive?"), "DEADLINE_EXCEEDED"),
                            (subject.ExtractionError("PAGE_LIMIT"), "PAGE_LIMIT"),
                            (subject.ExtractionError("attacker-chosen-detail"), "EXTRACTION_FAILED")]:
            with self.subTest(code=code):
                result = subject.extract_pages(pdf_bytes(["text"]), ".pdf", runner=FakeTrustedRunner(error=error))
                self.assertEqual((result.status, result.error_code, result.pages), ("failed", code, ()))
                self.assertNotIn("private", repr(result))

    def test_missing_dependency_is_an_unavailable_parser_failure(self):
        original_import = __import__
        def blocked_import(name, *args, **kwargs):
            if name == "pypdf":
                raise ImportError("private dependency path")
            return original_import(name, *args, **kwargs)
        with patch("builtins.__import__", side_effect=blocked_import):
            self.assert_code("DEPENDENCY_UNAVAILABLE", b"%PDF-1.7\n%%EOF", ".pdf")


if __name__ == "__main__":
    unittest.main()
