"""Fail-closed, bytes-only courseware extraction foundation, with no callers.

``extract_pages`` is the only application-facing entry point. Its default is
UNAVAILABLE: this module supplies no sandbox/process implementation. A deployment
must inject a separately reviewed TrustedProcessRunner that proves OS isolation
and hard limits before untrusted bytes enter a parser. A boolean supplied by an
arbitrary caller is not sandbox evidence; runners are trusted application code,
never request data. This module does not certify any runner or enable rollout.

``extract_pages_untrusted`` is INTERNAL WORKER-ONLY. It applies logical input,
archive, XML, page and text budgets but is NOT a sandbox or a hard time/memory
bound. PDF decompression/library calls can consume resources before a logical
check. Never call it on a request thread, in a thread timeout, or as a fallback.

No paths, URLs, environment/configuration, credentials, network clients, OCR,
subprocess launch, shell, logging, persistence, or provider calls are used here.
An injected runner owns supervision, environment sanitization and kill/reap;
receiving a timeout exception alone is not evidence that a process was stopped.
"""
from __future__ import annotations

from dataclasses import dataclass, fields
from hashlib import sha256
from io import BytesIO
import stat
from typing import Literal, Protocol
from xml.parsers import expat
from zipfile import ZipFile


_LIMIT_CAPS = {
    "max_file_bytes": 10 * 1024 * 1024,
    "max_pages": 200,
    "max_page_chars": 20_000,
    "max_total_chars": 200_000,
    "max_output_bytes": 1024 * 1024,
    "max_zip_entries": 2048,
    "max_zip_member_bytes": 16 * 1024 * 1024,
    "max_zip_total_bytes": 48 * 1024 * 1024,
    "max_zip_ratio": 100,
    "max_xml_depth": 64,
    "max_xml_elements": 100_000,
    "max_wall_seconds": 5,
    "max_cpu_seconds": 3,
    "max_memory_bytes": 256 * 1024 * 1024,
}


@dataclass(frozen=True)
class ExtractionLimits:
    """Positive integer budgets; callers may tighten, never expand these caps.

    Wall/CPU/memory and transport output budgets are runner obligations, not
    guarantees of the internal pure parser. max_output_bytes also bounds the
    aggregate UTF-8 text accepted from a runner. Actual IPC framing must fit the
    runner's total transport budget independently.
    """
    max_file_bytes: int = _LIMIT_CAPS["max_file_bytes"]
    max_pages: int = _LIMIT_CAPS["max_pages"]
    max_page_chars: int = _LIMIT_CAPS["max_page_chars"]
    max_total_chars: int = _LIMIT_CAPS["max_total_chars"]
    max_output_bytes: int = _LIMIT_CAPS["max_output_bytes"]
    max_zip_entries: int = _LIMIT_CAPS["max_zip_entries"]
    max_zip_member_bytes: int = _LIMIT_CAPS["max_zip_member_bytes"]
    max_zip_total_bytes: int = _LIMIT_CAPS["max_zip_total_bytes"]
    max_zip_ratio: int = _LIMIT_CAPS["max_zip_ratio"]
    max_xml_depth: int = _LIMIT_CAPS["max_xml_depth"]
    max_xml_elements: int = _LIMIT_CAPS["max_xml_elements"]
    max_wall_seconds: int = _LIMIT_CAPS["max_wall_seconds"]
    max_cpu_seconds: int = _LIMIT_CAPS["max_cpu_seconds"]
    max_memory_bytes: int = _LIMIT_CAPS["max_memory_bytes"]

    def __post_init__(self) -> None:
        for field in fields(self):
            value = getattr(self, field.name)
            if type(value) is not int or not 1 <= value <= _LIMIT_CAPS[field.name]:
                raise ValueError("limits must be positive integers within conservative caps")


DEFAULT_LIMITS = ExtractionLimits()


@dataclass(frozen=True)
class ExtractedPage:
    """One physical PDF page or PPTX slide, including blanks; one-based."""
    page: int
    text: str


@dataclass(frozen=True)
class ExtractionResult:
    status: Literal["ready", "failed", "unavailable"]
    content_sha256: str
    pages: tuple[ExtractedPage, ...] = ()
    error_code: str | None = None


_ERROR_CODES = frozenset({
    "INVALID_INPUT", "INVALID_LIMITS", "UNSUPPORTED_FORMAT", "EMPTY_INPUT",
    "FILE_LIMIT", "FORMAT_MISMATCH", "CORRUPT_SOURCE", "ENCRYPTED_SOURCE",
    "NO_EXTRACTABLE_TEXT", "PAGE_LIMIT", "PAGE_TEXT_LIMIT", "TOTAL_TEXT_LIMIT",
    "OUTPUT_LIMIT", "ZIP_ENTRY_LIMIT", "ZIP_MEMBER_LIMIT", "ZIP_TOTAL_LIMIT",
    "ZIP_RATIO_LIMIT", "UNSAFE_ARCHIVE", "UNSAFE_XML", "XML_DEPTH_LIMIT",
    "XML_ELEMENT_LIMIT", "DEPENDENCY_UNAVAILABLE", "ISOLATION_UNAVAILABLE",
    "HARD_LIMITS_UNAVAILABLE", "INVALID_OUTPUT", "RUNNER_FAILED",
    "DEADLINE_EXCEEDED", "RESOURCE_LIMIT", "EXTRACTION_FAILED",
})
_UNAVAILABLE_CODES = frozenset({
    "DEPENDENCY_UNAVAILABLE", "ISOLATION_UNAVAILABLE", "HARD_LIMITS_UNAVAILABLE",
})


class ExtractionError(Exception):
    """Stable, sanitized reason only; never expose underlying parser messages."""
    def __init__(self, code: str) -> None:
        self.code = code if type(code) is str and code in _ERROR_CODES else "EXTRACTION_FAILED"
        super().__init__(self.code)


class TrustedProcessRunner(Protocol):
    """Required deployment contract, NOT an implemented/verified sandbox.

    verify_support must independently establish an isolated, killable PROCESS
    with enforced wall deadline, CPU/address-space limits and bounded input and
    output IPC. It must deny worker network access and arbitrary filesystem
    access, expose only required read-only runtime/parser assets, clear inherited
    environment/credentials, close inherited handles and deny process escape.
    Environment variables, policy booleans and a thread timeout do not prove it.

    run must invoke extract_pages_untrusted only inside that established sandbox.
    Before returning OR raising it must kill and reap the worker and descendants
    as necessary; output is rejected after deadline or hard-limit failure. It
    must not deserialize untrusted pickle or return partial pages on failure.
    A launch/platform/enforcement failure must raise ISOLATION_UNAVAILABLE or
    HARD_LIMITS_UNAVAILABLE, never run the parser outside isolation. Deadline and
    resource failures use TimeoutError or ExtractionError('RESOURCE_LIMIT').

    Callers must not inject request-selected runners. An implementation requires
    independent integration/security tests; contract-double tests do not provide
    evidence of a real hard deadline, killed worker or effective confinement.
    """
    def verify_support(self, limits: ExtractionLimits) -> bool: ...

    def run(
        self, data: bytes, extension: str, limits: ExtractionLimits,
    ) -> tuple[ExtractedPage, ...]: ...


def _validate_input(data: bytes, extension: str, limits: ExtractionLimits) -> str:
    if type(limits) is not ExtractionLimits:
        raise ExtractionError("INVALID_LIMITS")
    # Recheck the immutable value object's contract at the untrusted boundary.
    try:
        limits.__post_init__()
    except (ValueError, TypeError, AttributeError):
        raise ExtractionError("INVALID_LIMITS") from None
    if type(data) is not bytes:
        raise ExtractionError("INVALID_INPUT")
    if not data:
        raise ExtractionError("EMPTY_INPUT")
    if len(data) > limits.max_file_bytes:
        raise ExtractionError("FILE_LIMIT")
    if type(extension) is not str or extension.casefold() not in (".pdf", ".pptx"):
        raise ExtractionError("UNSUPPORTED_FORMAT")
    extension = extension.casefold()
    if not data.startswith(b"%PDF-" if extension == ".pdf" else b"PK\x03\x04"):
        raise ExtractionError("FORMAT_MISMATCH")
    return extension


def _validate_pages(
    pages: tuple[ExtractedPage, ...], limits: ExtractionLimits, *, require_text: bool = True,
) -> None:
    if type(pages) is not tuple or not pages:
        raise ExtractionError("INVALID_OUTPUT")
    if len(pages) > limits.max_pages:
        raise ExtractionError("PAGE_LIMIT")
    total_chars = total_bytes = 0
    searchable = False
    for index, page in enumerate(pages, 1):
        # Exact-class instances can still be incomplete after a defective IPC
        # adapter; validate fields before reading them rather than leaking an
        # AttributeError outside the application-facing result boundary.
        if (type(page) is not ExtractedPage or type(getattr(page, "page", None)) is not int or
                page.page != index or type(getattr(page, "text", None)) is not str):
            raise ExtractionError("INVALID_OUTPUT")
        if len(page.text) > limits.max_page_chars:
            raise ExtractionError("PAGE_TEXT_LIMIT")
        total_chars += len(page.text)
        if total_chars > limits.max_total_chars:
            raise ExtractionError("TOTAL_TEXT_LIMIT")
        try:
            total_bytes += len(page.text.encode("utf-8", errors="strict"))
        except UnicodeError:
            raise ExtractionError("INVALID_OUTPUT") from None
        if total_bytes > limits.max_output_bytes:
            raise ExtractionError("OUTPUT_LIMIT")
        searchable = searchable or bool(page.text.strip())
    if require_text and not searchable:
        raise ExtractionError("NO_EXTRACTABLE_TEXT")


def _preflight_pptx(data: bytes, limits: ExtractionLimits) -> None:
    """Inspect every ZIP member and XML declaration, without filesystem writes."""
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > limits.max_zip_entries:
            raise ExtractionError("ZIP_ENTRY_LIMIT")
        names: set[str] = set()
        total = element_count = 0
        for entry in entries:
            name = entry.filename
            components = name.rstrip("/").split("/")
            if (name in names or name.startswith("/") or "\\" in name or
                    any(part in ("", ".", "..") for part in components) or
                    entry.orig_filename != name or ":" in name or
                    stat.S_ISLNK(entry.external_attr >> 16)):
                raise ExtractionError("UNSAFE_ARCHIVE")
            names.add(name)
            if entry.flag_bits & 1:
                raise ExtractionError("ENCRYPTED_SOURCE")
            if entry.file_size > limits.max_zip_member_bytes:
                raise ExtractionError("ZIP_MEMBER_LIMIT")
            total += entry.file_size
            if total > limits.max_zip_total_bytes:
                raise ExtractionError("ZIP_TOTAL_LIMIT")
            if entry.file_size > max(1, entry.compress_size) * limits.max_zip_ratio:
                raise ExtractionError("ZIP_RATIO_LIMIT")

        xml_overrides: set[str] = set()
        xml_extensions: set[str] = set()
        # Read the content-type manifest first: OOXML parts can legally use a
        # non-.xml filename. Extension-only checks would miss hostile XML there.
        entries.sort(key=lambda item: item.filename != "[Content_Types].xml")
        for entry in entries:
            name = entry.filename
            is_manifest = name == "[Content_Types].xml"
            is_xml = (name.casefold().endswith((".xml", ".rels")) or
                      name in xml_overrides or name.rsplit(".", 1)[-1].casefold() in xml_extensions)
            parser = expat.ParserCreate(namespace_separator="}") if is_xml else None
            depth = 0
            if parser is not None:
                def reject_xml(*args):
                    raise ExtractionError("UNSAFE_XML")

                def start_element(name, attributes):
                    nonlocal depth, element_count
                    depth += 1
                    element_count += 1
                    if depth > limits.max_xml_depth:
                        raise ExtractionError("XML_DEPTH_LIMIT")
                    if element_count > limits.max_xml_elements:
                        raise ExtractionError("XML_ELEMENT_LIMIT")
                    content_type = attributes.get("ContentType", "").casefold()
                    if is_manifest and (content_type.endswith("+xml") or content_type in ("application/xml", "text/xml")):
                        prefix = "http://schemas.openxmlformats.org/package/2006/content-types}"
                        if name == prefix + "Override":
                            xml_overrides.add(attributes.get("PartName", "").lstrip("/"))
                        elif name == prefix + "Default":
                            xml_extensions.add(attributes.get("Extension", "").casefold())

                def end_element(name):
                    nonlocal depth
                    depth -= 1

                parser.StartDoctypeDeclHandler = reject_xml
                parser.EntityDeclHandler = reject_xml
                parser.ExternalEntityRefHandler = reject_xml
                parser.StartElementHandler = start_element
                parser.EndElementHandler = end_element
                parser.SetParamEntityParsing(expat.XML_PARAM_ENTITY_PARSING_NEVER)
            expanded = 0
            with archive.open(entry) as source:
                while chunk := source.read(min(65536, limits.max_zip_member_bytes + 1 - expanded)):
                    expanded += len(chunk)
                    if expanded > limits.max_zip_member_bytes:
                        raise ExtractionError("ZIP_MEMBER_LIMIT")
                    if parser is not None:
                        parser.Parse(chunk, False)
            if expanded != entry.file_size:
                raise ExtractionError("CORRUPT_SOURCE")
            if parser is not None:
                parser.Parse(b"", True)
        if not {"[Content_Types].xml", "_rels/.rels", "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels"}.issubset(names):
            raise ExtractionError("CORRUPT_SOURCE")


def _slide_text(shapes, limits: ExtractionLimits) -> str:
    fragments: list[str] = []
    size = 0
    stack = [iter(shapes)]
    while stack:
        try:
            shape = next(stack[-1])
        except StopIteration:
            stack.pop()
            continue
        texts = []
        if shape.has_text_frame:
            texts.append(shape.text_frame.text)
        if shape.has_table:
            texts.extend(cell.text for row in shape.table.rows for cell in row.cells if not cell.is_spanned)
        for text in texts:
            size += len(text) + (1 if fragments else 0)
            if size > limits.max_page_chars:
                raise ExtractionError("PAGE_TEXT_LIMIT")
            fragments.append(text)
        if hasattr(shape, "shapes"):
            if len(stack) >= limits.max_xml_depth:
                raise ExtractionError("XML_DEPTH_LIMIT")
            stack.append(iter(shape.shapes))
    return "\n".join(fragments)


def extract_pages_untrusted(
    data: bytes, extension: str, limits: ExtractionLimits = DEFAULT_LIMITS,
) -> tuple[ExtractedPage, ...]:
    """INTERNAL ONLY: logical-bounded parser; requires a trusted sandbox process.

    Never fabricates page numbers, truncates text/pages to fit, decrypts files,
    follows external links, OCRs images, or returns partial success on an error.
    Blanks remain in the tuple so later page/slide indices retain their meaning.
    """
    extension = _validate_input(data, extension, limits)
    try:
        if extension == ".pdf":
            from pypdf import PdfReader
            document = PdfReader(BytesIO(data), strict=True)
            if document.is_encrypted:
                raise ExtractionError("ENCRYPTED_SOURCE")
            source_pages = document.pages
        else:
            _preflight_pptx(data, limits)
            from pptx import Presentation
            document = Presentation(BytesIO(data))
            source_pages = document.slides
        if len(source_pages) > limits.max_pages:
            raise ExtractionError("PAGE_LIMIT")
        pages: list[ExtractedPage] = []
        for index, page in enumerate(source_pages, 1):
            text = page.extract_text() if extension == ".pdf" else _slide_text(page.shapes, limits)
            if type(text) is not str:
                raise ExtractionError("CORRUPT_SOURCE")
            pages.append(ExtractedPage(index, text))
            # Aggregate budgets apply to blanks too; only the final result
            # requires extractable text. Never parse another page after overflow.
            _validate_pages(tuple(pages), limits, require_text=False)
        result = tuple(pages)
        _validate_pages(result, limits)
        return result
    except ExtractionError:
        raise
    except ImportError:
        raise ExtractionError("DEPENDENCY_UNAVAILABLE") from None
    except MemoryError:
        raise ExtractionError("RESOURCE_LIMIT") from None
    except Exception:
        raise ExtractionError("CORRUPT_SOURCE") from None


def extract_pages(
    data: bytes, extension: str, limits: ExtractionLimits = DEFAULT_LIMITS, *,
    runner: TrustedProcessRunner | None = None,
) -> ExtractionResult:
    """Fail closed unless trusted deployment code supplies proven supervision.

    SHA-256 is of exact input bytes, including blanks/metadata. It is omitted
    for non-bytes or inputs over the module's absolute file cap to avoid hashing
    unbounded input. Error results never include any extracted/partial pages.
    """
    digest = sha256(data).hexdigest() if type(data) is bytes and len(data) <= _LIMIT_CAPS["max_file_bytes"] else ""
    try:
        extension = _validate_input(data, extension, limits)
        if runner is None:
            raise ExtractionError("ISOLATION_UNAVAILABLE")
        try:
            supported = runner.verify_support(limits)
        except Exception:
            raise ExtractionError("ISOLATION_UNAVAILABLE") from None
        if supported is not True:
            raise ExtractionError("ISOLATION_UNAVAILABLE")
        try:
            pages = runner.run(data, extension, limits)
        except ExtractionError:
            raise
        except TimeoutError:
            raise ExtractionError("DEADLINE_EXCEEDED") from None
        except MemoryError:
            raise ExtractionError("RESOURCE_LIMIT") from None
        except Exception:
            raise ExtractionError("RUNNER_FAILED") from None
        _validate_pages(pages, limits)
        return ExtractionResult("ready", digest, pages)
    except ExtractionError as error:
        # Runner exceptions are mutable even when their constructor validated
        # them. Re-allowlist at the final boundary, including missing/unhashable
        # codes, without returning the original exception message or details.
        code = getattr(error, "code", None)
        if type(code) is not str or code not in _ERROR_CODES:
            code = "EXTRACTION_FAILED"
        status = "unavailable" if code in _UNAVAILABLE_CODES else "failed"
        return ExtractionResult(status, digest, error_code=code)
