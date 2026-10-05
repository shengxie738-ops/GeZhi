"""Detached deterministic native-text PPTX builder, with no integration effects."""
from __future__ import annotations

from io import BytesIO
import json
from time import monotonic

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from app.schemas.teacher_work import SlideModel
from app.services.teacher_work.exporters.theme import (
    ACCENT, BACKGROUND, BODY_FONT_SIZE, EMU_PER_POINT, HEIGHT, INK,
    PPTX_FONT, PPTX_FONT_DISCLOSURE, SECONDARY, TEMPLATE_VERSION,
    TITLE_FONT_SIZE, WIDTH, ThemeVersion, check_density, text_boxes,
)
from app.services.teacher_work.exporters.validation import (
    _check_stage_time, _finalize_generated_office_bytes, _require_xml_text,
)


def _east_asian(properties) -> None:
    element = properties.find(qn("a:ea"))
    if element is None:
        element = OxmlElement("a:ea")
        properties.append(element)
    element.set("typeface", PPTX_FONT)


def _paragraph(paragraph, text: str, size: int, color: str, bold: bool = False) -> None:
    # A soft break within an item is a:br; item boundaries remain separate a:p.
    paragraph.text = text.replace("\n", "\v")
    paragraph.font.name = PPTX_FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    paragraph.line_spacing = 1.2
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(6)
    _east_asian(paragraph._p.get_or_add_pPr().get_or_add_defRPr())
    for run in paragraph.runs:
        run.font.name = PPTX_FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        _east_asian(run._r.get_or_add_rPr())


def _box(slide, name: str, geometry, items, size: int, color: str, bold=False):
    shape = slide.shapes.add_textbox(*(Pt(value) for value in geometry))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Pt(8)
    frame.margin_top = frame.margin_bottom = Pt(8 if size >= 20 else 0)
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for index, item in enumerate(items or ("",)):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _paragraph(paragraph, item, size, color, bold)
    return shape


def build_pptx(model: SlideModel, theme: ThemeVersion) -> bytes:
    """Return actual editable Office bytes, or a controlled layout/bounds failure."""
    started = monotonic()
    if type(model) is not SlideModel or type(theme) is not ThemeVersion or theme.version != TEMPLATE_VERSION:
        raise ValueError("INVALID_EXPORT_INPUT")
    model = SlideModel.model_validate(model.model_dump())
    _require_xml_text(model.model_dump())
    # Preflight every page before constructing a single slide.
    for page, snapshot in enumerate(model.slides, 1):
        boxes = text_boxes(snapshot.layout)
        check_density((snapshot.title,), boxes["gezhi:title"], TITLE_FONT_SIZE, page)
        check_density((snapshot.source_note,), boxes["gezhi:source"], 10, page)
        if snapshot.layout == "two_column":
            for name, items in zip(("gezhi:left", "gezhi:right"), snapshot.columns):
                check_density(items, boxes[name], BODY_FONT_SIZE, page)
        else:
            check_density(snapshot.body, boxes["gezhi:body"], BODY_FONT_SIZE, page)
        _check_stage_time(started)
    presentation = Presentation()
    presentation.slide_width = WIDTH * EMU_PER_POINT
    presentation.slide_height = HEIGHT * EMU_PER_POINT
    presentation.core_properties.author = "GeZhi"
    presentation.core_properties.last_modified_by = "GeZhi"
    presentation.core_properties.title = "教师私人备课资料"
    for snapshot in model.slides:
        _check_stage_time(started)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(BACKGROUND)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(40), Pt(128), Pt(880), Pt(3))
        accent.name = "gezhi-accent-rule"
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
        accent.line.fill.background()
        boxes = text_boxes(snapshot.layout)
        title = _box(slide, "gezhi:title", boxes["gezhi:title"], (snapshot.title,), TITLE_FONT_SIZE, INK, True)
        title._element.nvSpPr.cNvPr.set("descr", json.dumps({
            "layout": snapshot.layout, "evidence_refs": [str(ref) for ref in snapshot.evidence_refs],
        }, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
        if snapshot.layout == "two_column":
            for name, items in zip(("gezhi:left", "gezhi:right"), snapshot.columns):
                _box(slide, name, boxes[name], items, BODY_FONT_SIZE, INK)
        else:
            _box(slide, "gezhi:body", boxes["gezhi:body"], snapshot.body, BODY_FONT_SIZE, INK)
        _box(slide, "gezhi:source", boxes["gezhi:source"], (snapshot.source_note,), 10, SECONDARY)
        _box(slide, "gezhi:font", boxes["gezhi:font"], (PPTX_FONT_DISCLOSURE,), 8, SECONDARY)
        # Creation is intentional here; validation must never create missing notes.
        notes = slide.notes_slide.notes_text_frame
        if notes is None:
            raise ValueError("PPTX_NOTES_UNAVAILABLE")
        notes.text = snapshot.notes.replace("\n", "\v")
        for paragraph in notes.paragraphs:
            _paragraph(paragraph, paragraph.text.replace("\v", "\n"), 12, SECONDARY)
    output = BytesIO()
    presentation.save(output)
    _check_stage_time(started)
    raw = _finalize_generated_office_bytes("pptx", output.getvalue())
    _check_stage_time(started)
    return raw
