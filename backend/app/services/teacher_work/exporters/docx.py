"""Strict frozen Work mapping around the unchanged legacy DOCX renderer body."""
from __future__ import annotations

from io import BytesIO
from time import monotonic

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from app.schemas.teacher_work import LessonSnapshot
from app.services.teacher_lesson_prep.docx_exporter import build_lesson_plan_docx
from app.services.teacher_work.exporters.theme import DOCX_FONT_DISCLOSURE
from app.services.teacher_work.exporters.validation import (
    _check_stage_time, _docx_expected_paragraphs, _finalize_generated_office_bytes,
    _parts_and_roots, _paragraphs, _require_xml_text,
    check_docx_snapshot_text,
)


def _exact_legacy_text(snapshot: LessonSnapshot) -> None:
    """Reject trimming/dropping/coercion instead of modifying the frozen lesson."""
    check_docx_snapshot_text(snapshot)


def build_docx(snapshot: LessonSnapshot) -> bytes:
    started = monotonic()
    if type(snapshot) is not LessonSnapshot:
        raise ValueError("INVALID_EXPORT_INPUT")
    snapshot = LessonSnapshot.model_validate(snapshot.model_dump())
    _exact_legacy_text(snapshot)
    # Full frozen JSON projection uses the exact legacy field names and plain lists/dicts.
    raw = build_lesson_plan_docx(snapshot.model_dump(mode="json"))
    _check_stage_time(started)
    raw = _finalize_generated_office_bytes("docx", raw)
    document = Document(BytesIO(raw))
    paragraph = document.add_paragraph()
    run = paragraph.add_run(DOCX_FONT_DISCLOSURE)
    run.font.name = "宋体"
    run.font.size = Pt(10)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    document.core_properties.author = "GeZhi"
    document.core_properties.last_modified_by = "GeZhi"
    document.core_properties.title = snapshot.title
    output = BytesIO()
    document.save(output)
    _check_stage_time(started)
    raw = _finalize_generated_office_bytes("docx", output.getvalue())
    _, roots = _parts_and_roots("docx", raw)
    if _paragraphs(roots["word/document.xml"], "w") != _docx_expected_paragraphs(snapshot):
        raise ValueError("DOCX_UNREPRESENTABLE")
    _check_stage_time(started)
    return raw
