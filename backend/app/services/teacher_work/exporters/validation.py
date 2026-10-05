"""Bounded, read-only OOXML validation and exact trusted-bootstrap finalization.

No extraction, path/URL readers, repair, resolver, XSLT, image decode or font APIs.
Convenience readers run only after the independent ZIP/XML/projection preflight.
"""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
import posixpath
import re
from time import monotonic
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo
from zlib import error as DeflateError

from app.schemas.teacher_work import PackageVersionDTO, ValidationSummary
from app.services.teacher_work.types import canonical_digest
from app.services.teacher_work.exporters.theme import (
    BODY_FONT_SIZE, DOCX_EXPORTER_VERSION, DOCX_FONT_DISCLOSURE, EMU_PER_POINT,
    HEIGHT, PPTX_EXPORTER_VERSION, PPTX_FONT, PPTX_FONT_DISCLOSURE,
    TEMPLATE_VERSION, WIDTH, check_density, text_boxes,
)


FILE_CAP = 10 * 1024 * 1024
MEMBER_CAP = 256
XML_CAP = 1024 * 1024
OTHER_CAP = 128 * 1024
TOTAL_CAP = 8 * 1024 * 1024
INFLATION_CAP = 1000
# DEFLATE option bits1/2, data descriptor bit3 and UTF-8 bit11 only.
# The trusted PPTX template has flags value6; unsupported bits5/6 are32/64.
ALLOWED_ZIP_FLAGS = (1 << 1) | (1 << 2) | (1 << 3) | (1 << 11)
STAGE_SECONDS = 30
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
}
REL = NS["r"] + "/"
PACKAGE_REL = NS["rel"] + "/"
PPT_TYPE = "application/vnd.openxmlformats-officedocument.presentationml."
WORD_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml."
CORE_TYPE = "application/vnd.openxmlformats-package.core-properties+xml"
APP_TYPE = "application/vnd.openxmlformats-officedocument.extended-properties+xml"
THEME_TYPE = "application/vnd.openxmlformats-officedocument.theme+xml"
CUSTOM_TYPE = "application/vnd.openxmlformats-officedocument.customXmlProperties+xml"
PRINTER_TYPE = PPT_TYPE + "printerSettings"
REL_TYPE = "application/vnd.openxmlformats-package.relationships+xml"
THUMBNAIL = "docProps/thumbnail.jpeg"
PRINTER = "ppt/printerSettings/printerSettings1.bin"
BIBLIOGRAPHY = {
    "customXml/item1.xml": "a86086ffc5d8e83ebd6c71a55d1d2efaa31b137977f5f3a752366e1023612144",
    "customXml/itemProps1.xml": "c542307b13ec29a8b546217bb37936ab4822e044b265d2952985ec3d6afed24e",
    "customXml/_rels/item1.xml.rels": "32fa6ddf5bdffdeec97e77d0805928f4d2411f23a3335d1ee0449b740ce1d84d",
}
THUMB_HASH = {
    "pptx": "116f9aa8ea038e39ef47ee99a23592ab569cf77feafecd0a30e222369e931d25",
    "docx": "96367138dc44ce09bf2c8f0f8e49348a1478d2c5c0af69bbc2bbc38b63cdcead",
}
PRINTER_HASH = "d7768f87e07d29634782e448ece1cddd05a52b9499254222b1190bc8f6dc579e"
MAIN = {"pptx": "ppt/presentation.xml", "docx": "word/document.xml"}
WARNING = "仅结构与内容校验；实际 Office 全页编辑、排版与字体替代仍待核对"


def _check_stage_time(started: float) -> None:
    """Cooperative stage checkpoints; orchestration owns hard deadline/cancellation."""
    if monotonic() - started > STAGE_SECONDS:
        raise ValueError("OFFICE_STAGE_TIMEOUT")


def _require_xml_text(value, error="INVALID_XML_TEXT") -> None:
    if isinstance(value, str):
        for char in value:
            number = ord(char)
            if not (number in (9, 10) or 32 <= number <= 0xD7FF or 0xE000 <= number <= 0xFFFD or 0x10000 <= number <= 0x10FFFF):
                raise ValueError(error)
    elif isinstance(value, dict):
        for item in value.values():
            _require_xml_text(item, error)
    elif isinstance(value, (tuple, list)):
        for item in value:
            _require_xml_text(item, error)


def _normalized_name(name: str) -> str:
    if not name or name.startswith("/") or "\\" in name or ":" in name or any(ord(c) < 32 for c in name):
        raise ValueError("UNSAFE_PART_NAME")
    normalized = posixpath.normpath(name)
    if normalized != name or normalized in (".", "..") or normalized.startswith("../"):
        raise ValueError("UNSAFE_PART_NAME")
    return normalized


def _xml(raw: bytes):
    # Decode BOM inputs before rejecting so obfuscated DTDs never reach a native
    # reader. Successful generated parts are UTF-8 only; other encodings fail
    # closed. No arbitrary codec lookup, resolver, DTD or entity processing.
    non_utf8 = False
    if raw.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        text = raw.decode("utf-32")
        non_utf8 = True
    elif raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        text = raw.decode("utf-16")
        non_utf8 = True
    else:
        text = raw.decode("utf-8")
        if text.startswith("\ufeff"):
            text = text[1:]
    if "<!--" in text:
        raise ValueError("UNEXPECTED_XML_COMMENT")
    if re.search(r"<!\s*(?:DOCTYPE|ENTITY)\b", text, re.IGNORECASE):
        raise ValueError("UNSAFE_XML_DECLARATION")
    if re.search(r"<\?(?!xml\s)", text, re.IGNORECASE):
        raise ValueError("UNSAFE_XML_PROCESSING_INSTRUCTION")
    if non_utf8 or "\x00" in text:
        raise ValueError("UNSUPPORTED_XML_ENCODING")
    declaration = re.match(r"\s*<\?xml\s+[^?]*encoding\s*=\s*['\"]([^'\"]+)", text, re.IGNORECASE)
    if declaration and declaration.group(1).lower().replace("_", "-") != "utf-8":
        raise ValueError("UNSUPPORTED_XML_ENCODING")
    return ET.fromstring(text)


def _read_parts(raw: bytes) -> dict[str, bytes]:
    if type(raw) is not bytes or not raw or len(raw) > FILE_CAP:
        raise ValueError("OFFICE_BYTES_LIMIT")
    parts = {}
    total = 0
    with ZipFile(BytesIO(raw)) as package:
        infos = package.infolist()
        if package.comment:
            raise ValueError("UNEXPECTED_ZIP_METADATA")
        if not infos or len(infos) > MEMBER_CAP:
            raise ValueError("OFFICE_MEMBER_LIMIT")
        for info in infos:
            name = _normalized_name(info.filename)
            if info.extra or info.comment:
                raise ValueError("UNEXPECTED_ZIP_METADATA")
            if name in parts or info.is_dir() or info.flag_bits & ~ALLOWED_ZIP_FLAGS or info.compress_type not in (ZIP_STORED, ZIP_DEFLATED):
                raise ValueError("UNSAFE_ZIP_MEMBER")
            limit = XML_CAP if name.endswith((".xml", ".rels")) else OTHER_CAP
            # [Content_Types].xml is XML as well. All caps are checked before reading.
            if info.file_size > limit or info.file_size < 0 or info.compress_size < 0:
                raise ValueError("OFFICE_PART_LIMIT")
            if info.file_size > max(1, info.compress_size) * INFLATION_CAP:
                raise ValueError("OFFICE_INFLATION_LIMIT")
            total += info.file_size
            if total > TOTAL_CAP:
                raise ValueError("OFFICE_TOTAL_LIMIT")
            header = raw[info.header_offset:info.header_offset + 30]
            if len(header) != 30 or header[:4] != b"PK\x03\x04" or int.from_bytes(header[6:8], "little") != info.flag_bits or int.from_bytes(header[8:10], "little") != info.compress_type:
                raise ValueError("INCONSISTENT_ZIP_HEADER")
            with package.open(info) as stream:
                content = stream.read(limit + 1)
            if len(content) != info.file_size or len(content) > limit:
                raise ValueError("OFFICE_PART_LIMIT")
            parts[name] = content
    return parts


def _part_type(kind: str, name: str, bootstrap: bool) -> str:
    if name == "[Content_Types].xml":
        return ""
    if name.endswith(".rels"):
        # Owner path is verified independently, including deleted bibliography.
        return REL_TYPE
    if name == "docProps/core.xml":
        return CORE_TYPE
    if name == "docProps/app.xml":
        return APP_TYPE
    if bootstrap:
        if name == THUMBNAIL:
            return "image/jpeg"
        if kind == "pptx" and name == PRINTER:
            return PRINTER_TYPE
        if kind == "docx" and name in BIBLIOGRAPHY:
            return CUSTOM_TYPE if name.endswith("itemProps1.xml") else "application/xml"
    if kind == "pptx":
        simple = {
            "ppt/presentation.xml": "presentation.main+xml", "ppt/presProps.xml": "presProps+xml",
            "ppt/viewProps.xml": "viewProps+xml", "ppt/tableStyles.xml": "tableStyles+xml",
        }
        if name in simple:
            return PPT_TYPE + simple[name]
        patterns = (
            (r"ppt/slides/slide(?:[1-9]|1[0-2])\.xml", "slide+xml"),
            (r"ppt/slideLayouts/slideLayout(?:[1-9]|1[01])\.xml", "slideLayout+xml"),
            (r"ppt/slideMasters/slideMaster1\.xml", "slideMaster+xml"),
            (r"ppt/notesSlides/notesSlide(?:[1-9]|1[0-2])\.xml", "notesSlide+xml"),
            (r"ppt/notesMasters/notesMaster1\.xml", "notesMaster+xml"),
        )
        for pattern, suffix in patterns:
            if re.fullmatch(pattern, name):
                return PPT_TYPE + suffix
        if re.fullmatch(r"ppt/theme/theme[12]\.xml", name):
            return THEME_TYPE
    else:
        simple = {
            "word/document.xml": WORD_TYPE + "document.main+xml",
            "word/numbering.xml": WORD_TYPE + "numbering+xml",
            "word/styles.xml": WORD_TYPE + "styles+xml",
            "word/stylesWithEffects.xml": "application/vnd.ms-word.stylesWithEffects+xml",
            "word/settings.xml": WORD_TYPE + "settings+xml",
            "word/webSettings.xml": WORD_TYPE + "webSettings+xml",
            "word/fontTable.xml": WORD_TYPE + "fontTable+xml",
            "word/theme/theme1.xml": THEME_TYPE,
        }
        if name in simple:
            return simple[name]
    raise ValueError("UNEXPECTED_OFFICE_PART")


def _expected_root(name: str):
    fixed = {
        "docProps/core.xml": "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}coreProperties",
        "docProps/app.xml": "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Properties",
        "ppt/presentation.xml": "p:presentation", "ppt/presProps.xml": "p:presentationPr",
        "ppt/viewProps.xml": "p:viewPr", "ppt/tableStyles.xml": "a:tblStyleLst",
        "word/document.xml": "w:document", "word/styles.xml": "w:styles",
        "word/stylesWithEffects.xml": "w:styles", "word/settings.xml": "w:settings",
        "word/webSettings.xml": "w:webSettings", "word/fontTable.xml": "w:fonts",
        "word/numbering.xml": "w:numbering",
    }
    expected = fixed.get(name)
    if expected is None:
        for prefix, tag in (("ppt/slides/", "p:sld"), ("ppt/slideLayouts/", "p:sldLayout"),
                            ("ppt/slideMasters/", "p:sldMaster"), ("ppt/notesSlides/", "p:notes"),
                            ("ppt/notesMasters/", "p:notesMaster"), ("ppt/theme/", "a:theme"),
                            ("word/theme/", "a:theme")):
            if name.startswith(prefix) and not name.endswith(".rels"):
                expected = tag
                break
    if expected and not expected.startswith("{"):
        prefix, local = expected.split(":")
        expected = "{" + NS[prefix] + "}" + local
    return expected


def _rel_owner(name: str) -> str:
    if name == "_rels/.rels":
        return ""
    if "/_rels/" not in name or not name.endswith(".rels"):
        raise ValueError("INVALID_RELATIONSHIP_PART")
    parent, local = name.rsplit("/_rels/", 1)
    return parent + "/" + local[:-5]


def _rel_name(owner: str) -> str:
    return "_rels/.rels" if not owner else posixpath.join(posixpath.dirname(owner), "_rels", posixpath.basename(owner) + ".rels")


def _target(owner: str, target: str) -> str:
    if not target or "\\" in target or ":" in target or "?" in target or "#" in target or "%" in target or target.startswith("/"):
        raise ValueError("UNSAFE_RELATIONSHIP_TARGET")
    resolved = posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
    return _normalized_name(resolved)


def _graph(kind: str, parts: dict[str, bytes], roots: dict, bootstrap=False) -> dict:
    if kind not in MAIN or "[Content_Types].xml" not in roots or "_rels/.rels" not in roots or MAIN[kind] not in roots:
        raise ValueError("MISSING_REQUIRED_PART")
    declarations = roots["[Content_Types].xml"]
    if declarations.tag != "{" + NS["ct"] + "}Types" or declarations.attrib or (declarations.text or "").strip():
        raise ValueError("INVALID_CONTENT_TYPES")
    defaults, overrides = {}, {}
    for element in declarations:
        if len(element) or (element.text or "").strip() or (element.tail or "").strip():
            raise ValueError("INVALID_CONTENT_TYPES")
        if element.tag == "{" + NS["ct"] + "}Default":
            if set(element.attrib) != {"Extension", "ContentType"}:
                raise ValueError("INVALID_CONTENT_TYPES")
            extension, content_type = element.attrib["Extension"], element.attrib["ContentType"]
            permitted = {"xml": "application/xml", "rels": REL_TYPE}
            if bootstrap:
                permitted["jpeg"] = "image/jpeg"
                if kind == "pptx":
                    permitted["bin"] = PRINTER_TYPE
            if extension in defaults or permitted.get(extension) != content_type:
                raise ValueError("INVALID_CONTENT_TYPES")
            defaults[extension] = content_type
        elif element.tag == "{" + NS["ct"] + "}Override":
            if set(element.attrib) != {"PartName", "ContentType"}:
                raise ValueError("INVALID_CONTENT_TYPES")
            name = element.attrib["PartName"]
            if not name.startswith("/"):
                raise ValueError("INVALID_CONTENT_TYPES")
            name = _normalized_name(name[1:])
            if name not in parts or name in overrides or _part_type(kind, name, bootstrap) != element.attrib["ContentType"]:
                raise ValueError("INVALID_CONTENT_TYPES")
            overrides[name] = element.attrib["ContentType"]
        else:
            raise ValueError("INVALID_CONTENT_TYPES")
    for name in parts:
        expected = _part_type(kind, name, bootstrap)
        expected_root = _expected_root(name)
        if expected_root and (name not in roots or roots[name].tag != expected_root):
            raise ValueError("INVALID_PART_ROOT")
        if name == "[Content_Types].xml":
            continue
        actual = overrides.get(name, defaults.get(name.rsplit(".", 1)[-1]))
        if actual != expected:
            raise ValueError("INVALID_CONTENT_TYPES")
    allowed = {REL + term for term in (
        "officeDocument", "extended-properties", "slide", "slideLayout", "slideMaster", "theme",
        "notesSlide", "notesMaster", "presProps", "viewProps", "tableStyles", "styles",
        "settings", "webSettings", "fontTable", "numbering",
    )} | {PACKAGE_REL + "metadata/core-properties", "http://schemas.microsoft.com/office/2007/relationships/stylesWithEffects"}
    if bootstrap:
        allowed.add(PACKAGE_REL + "metadata/thumbnail")
        allowed.add(REL + ("printerSettings" if kind == "pptx" else "customXml"))
        if kind == "docx":
            allowed.add(REL + "customXmlProps")
    relationships = {}
    incoming = {name: [] for name in parts}
    for name, root in roots.items():
        if not name.endswith(".rels"):
            continue
        owner = _rel_owner(name)
        if owner and owner not in parts:
            raise ValueError("DANGLING_RELATIONSHIP_OWNER")
        if root.tag != "{" + NS["rel"] + "}Relationships" or root.attrib or (root.text or "").strip():
            raise ValueError("INVALID_RELATIONSHIP_PART")
        local = {}
        for element in root:
            if element.tag != "{" + NS["rel"] + "}Relationship" or set(element.attrib) - {"Id", "Type", "Target", "TargetMode"} or len(element) or (element.text or "").strip() or (element.tail or "").strip():
                raise ValueError("INVALID_RELATIONSHIP")
            identifier = element.attrib["Id"]
            if not identifier or identifier in local or element.attrib.get("TargetMode") not in (None, "Internal") or element.attrib["Type"] not in allowed:
                raise ValueError("UNSAFE_RELATIONSHIP")
            target = _target(owner, element.attrib["Target"])
            if target not in parts:
                raise ValueError("DANGLING_RELATIONSHIP")
            local[identifier] = (element, target)
            incoming[target].append((owner, element))
        relationships[owner] = local
    office_edges = [(edge, target) for edge, target in relationships.get("", {}).values() if edge.attrib["Type"] == REL + "officeDocument"]
    if len(office_edges) != 1 or office_edges[0][1] != MAIN[kind]:
        raise ValueError("INVALID_MAIN_PART")
    # Every non-metadata part must be reachable; unreferenced unsafe content is not ignored.
    reachable = {"[Content_Types].xml", "_rels/.rels"}
    pending = [""]
    visited = set()
    while pending:
        owner = pending.pop()
        if owner in visited:
            continue
        visited.add(owner)
        if owner:
            reachable.add(owner)
        if _rel_name(owner) in parts:
            reachable.add(_rel_name(owner))
        for _, target in relationships.get(owner, {}).values():
            pending.append(target)
    if reachable != set(parts):
        raise ValueError("UNREFERENCED_OFFICE_PART")
    forbidden = {"pic", "graphicFrame", "blip", "oleObj", "oleObject", "OLEObject", "object", "control", "altChunk",
                 "fldSimple", "fldChar", "instrText", "hyperlink", "hlinkClick", "hlinkHover", "timing", "transition", "externalLink",
                 "pict", "imagedata", "svgBlip", "frame", "frameset", "subDoc", "movie", "videoFile", "audioFile",
                 "wavAudioFile", "contentPart", "embeddedFont", "embeddedFontLst", "embedRegular", "embedBold"}
    for name, root in roots.items():
        if name.endswith(".rels") or name == "[Content_Types].xml":
            continue
        # These exporters intentionally produce group-free native sp children.
        # Reject unsupported tree children rather than overlooking grouped text.
        native_children = {"{" + NS["p"] + "}" + tag for tag in ("nvGrpSpPr", "grpSpPr", "sp")}
        for tree in root.iter("{" + NS["p"] + "}spTree"):
            if any(child.tag not in native_children for child in tree):
                raise ValueError("UNSUPPORTED_NATIVE_SHAPE")
        for element in root.iter():
            if element.tag.rsplit("}", 1)[-1] in forbidden or "action" in element.attrib:
                raise ValueError("UNSAFE_OFFICE_MARKUP")
            for key, identifier in element.attrib.items():
                if key.startswith("{" + NS["r"] + "}") and key.rsplit("}", 1)[-1] in {"id", "embed", "link"}:
                    if identifier not in relationships.get(name, {}):
                        raise ValueError("MISSING_REFERENCED_RELATIONSHIP")
    return {"relationships": relationships, "incoming": incoming, "defaults": defaults, "overrides": overrides}


def _parts_and_roots(kind: str, raw: bytes, bootstrap=False):
    parts = _read_parts(raw)
    roots = {name: _xml(content) for name, content in parts.items() if name.endswith((".xml", ".rels"))}
    _graph(kind, parts, roots, bootstrap)
    return parts, roots


def _edge_is(edge, identifier, relation_type, target) -> bool:
    return edge.attrib == {"Id": identifier, "Type": relation_type, "Target": target} and not len(edge) and not (edge.text or "").strip() and not (edge.tail or "").strip()


def _remove_cluster(parts, roots, graph, names, hashes, owner, identifier, relation_type, target):
    present = [name in parts for name in names]
    matching = [(local_owner, edge) for name in names for local_owner, edge in graph["incoming"].get(name, [])
                if local_owner not in names]
    if not any(present):
        if matching or any(edge.attrib.get("Type") == relation_type for edge, _ in graph["relationships"].get(owner, {}).values()):
            raise ValueError("PARTIAL_BOOTSTRAP_CLUSTER")
        return False
    if not all(present) or any(sha256(parts[name]).hexdigest() != digest for name, digest in hashes.items()):
        raise ValueError("CHANGED_BOOTSTRAP_PAYLOAD")
    if len(matching) != 1 or matching[0][0] != owner or not _edge_is(matching[0][1], identifier, relation_type, target):
        raise ValueError("CHANGED_BOOTSTRAP_RELATIONSHIP")
    # A deleted relationship must not be referenced by retained owner XML.
    if owner and any(value == identifier for element in roots[owner].iter() for key, value in element.attrib.items()
                     if key.startswith("{" + NS["r"] + "}")):
        raise ValueError("BOOTSTRAP_RELATIONSHIP_IN_USE")
    roots[_rel_name(owner)].remove(matching[0][1])
    for name in names:
        del parts[name]
        roots.pop(name, None)
    return True


def _finalize_generated_office_bytes(kind: str, raw: bytes) -> bytes:
    """Only exact reviewed optional bootstrap parts and timestamp metadata change.

    Complete absence is idempotent. Unexpected content is rejected, never cleaned
    into acceptance. Relationship attributes are owner-local and payload-pinned.
    """
    started = monotonic()
    parts, roots = _parts_and_roots(kind, raw, bootstrap=True)
    graph = _graph(kind, parts, roots, bootstrap=True)
    removed_thumb = _remove_cluster(parts, roots, graph, (THUMBNAIL,), {THUMBNAIL: THUMB_HASH[kind]},
        "", "rId2", PACKAGE_REL + "metadata/thumbnail", THUMBNAIL)
    removed_printer = False
    removed_bib = False
    if kind == "pptx":
        removed_printer = _remove_cluster(parts, roots, graph, (PRINTER,), {PRINTER: PRINTER_HASH},
            MAIN[kind], "rId2", REL + "printerSettings", "printerSettings/printerSettings1.bin")
    else:
        if any(name in parts for name in BIBLIOGRAPHY):
            child = roots.get("customXml/_rels/item1.xml.rels")
            if child is None or len(child) != 1 or not _edge_is(child[0], "rId1", REL + "customXmlProps", "itemProps1.xml"):
                raise ValueError("CHANGED_BOOTSTRAP_RELATIONSHIP")
        # Relationship serialization is upstream-generated; pin the data payloads
        # and exact owner-local child attributes rather than its XML whitespace.
        data_hashes = {name: digest for name, digest in BIBLIOGRAPHY.items() if not name.endswith(".rels")}
        removed_bib = _remove_cluster(parts, roots, graph, tuple(BIBLIOGRAPHY), data_hashes,
            MAIN[kind], "rId1", REL + "customXml", "../customXml/item1.xml")
    types = roots["[Content_Types].xml"]
    for element in list(types):
        if element.tag == "{" + NS["ct"] + "}Default" and element.attrib.get("Extension") in {"jpeg", "bin"}:
            extension = element.attrib["Extension"]
            removed = removed_thumb if extension == "jpeg" else removed_printer
            expected = "image/jpeg" if extension == "jpeg" else PRINTER_TYPE
            if not removed or element.attrib["ContentType"] != expected or any(name.endswith("." + extension) for name in parts):
                raise ValueError("UNEXPECTED_BOOTSTRAP_DECLARATION")
            types.remove(element)
        elif element.attrib.get("PartName") == "/customXml/itemProps1.xml":
            if not removed_bib or element.attrib.get("ContentType") != CUSTOM_TYPE:
                raise ValueError("UNEXPECTED_BOOTSTRAP_DECLARATION")
            types.remove(element)
    # Serialize only touched structural declarations/relationship XML; not document content.
    changed = {"[Content_Types].xml", "_rels/.rels"}
    if removed_printer or removed_bib:
        changed.add(_rel_name(MAIN[kind]))
    for name in changed:
        parts[name] = ET.tostring(roots[name], encoding="utf-8", xml_declaration=True)
    # Only these three core timestamp fields are normalized. No arbitrary stripping.
    core = roots.get("docProps/core.xml")
    if core is not None:
        # Preserve the QName namespace of the unchanged xsi:type attribute value.
        # ElementTree otherwise discards an explicit dcterms prefix declaration.
        core.set("xmlns:dcterms", "http://purl.org/dc/terms/")
        fields = {"{http://purl.org/dc/terms/}created", "{http://purl.org/dc/terms/}modified",
                  "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastPrinted"}
        for element in core:
            if element.tag in fields:
                element.text = "2000-01-01T00:00:00Z"
        parts["docProps/core.xml"] = ET.tostring(core, encoding="utf-8", xml_declaration=True)
    clean_roots = {name: _xml(content) for name, content in parts.items() if name.endswith((".xml", ".rels"))}
    _graph(kind, parts, clean_roots)
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as package:
        for name in sorted(parts):
            info = ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0x20
            package.writestr(info, parts[name])
            _check_stage_time(started)
    result = output.getvalue()
    _parts_and_roots(kind, result)
    _check_stage_time(started)
    return result


def _paragraph(element, dialect):
    result = []
    for child in element.iter():
        if child.tag == "{" + NS[dialect] + "}t":
            result.append(child.text or "")
        elif child.tag in {"{" + NS[dialect] + "}br", "{" + NS[dialect] + "}cr"}:
            result.append("\n")
        elif child.tag == "{" + NS[dialect] + "}tab":
            result.append("\t")
    return "".join(result)


def _paragraphs(element, dialect):
    return [_paragraph(paragraph, dialect) for paragraph in element.findall(".//" + dialect + ":p", NS)]


def _docx_expected_paragraphs(lesson):
    """Exact existing renderer projection, including paragraphs inside both tables."""
    values = [lesson.title, (lesson.course_name + " 教案") if lesson.course_name else "教案",
              "课程名称", lesson.course_name or "—", "授课对象", lesson.audience or "—",
              "备课主题", lesson.topic or "—", "课时", f"{lesson.duration_minutes} 分钟"]
    sections = []
    for field, heading in (("objectives", "教学目标"), ("key_points", "教学重点"), ("difficulties", "教学难点")):
        if getattr(lesson, field):
            sections.append((heading, [f"{index}. {text}" for index, text in enumerate(getattr(lesson, field), 1)]))
    flow = ["教学环节", "时长(分钟)", "教学内容"]
    for stage in lesson.teaching_flow:
        flow.extend((stage.stage, str(stage.minutes), stage.content))
    sections.append(("教学流程", flow))
    for field, heading in (("questions", "课堂提问"), ("exercises", "课堂练习"), ("homework", "课后作业")):
        if getattr(lesson, field):
            sections.append((heading, [f"{index}. {text}" for index, text in enumerate(getattr(lesson, field), 1)]))
    if lesson.summary:
        sections.append(("课件总结", lesson.summary.splitlines()))
    if lesson.citations:
        citations = []
        for index, citation in enumerate(lesson.citations, 1):
            citations.append(f"{index}. {citation.name or '未命名课件'}" + (f" 第{citation.page}页" if citation.page else ""))
            if citation.excerpt:
                citations.append("摘录：" + citation.excerpt)
        sections.append(("课件依据", citations))
    for index, (heading, content) in enumerate(sections):
        values.append("一二三四五六七八九十"[index] + "、" + heading)
        values.extend(content)
    return values + [DOCX_FONT_DISCLOSURE]


def _checked_version(version):
    if type(version) is not PackageVersionDTO:
        raise ValueError("INVALID_VERSION")
    version = PackageVersionDTO.model_validate(version.model_dump())
    frozen = version.model_dump(mode="json", include={"lesson", "slides", "source_snapshots"})
    _require_xml_text(frozen)
    if canonical_digest(frozen) != version.content_digest:
        raise ValueError("CONTENT_DIGEST_MISMATCH")
    if version.template_version != TEMPLATE_VERSION or tuple(version.exporter_versions) != (PPTX_EXPORTER_VERSION, DOCX_EXPORTER_VERSION):
        raise ValueError("EXPORTER_VERSION_MISMATCH")
    evidence = {item.evidence_id for item in version.source_snapshots}
    if len(evidence) != len(version.source_snapshots) or any(item.task_id != version.task_id for item in version.source_snapshots):
        raise ValueError("EVIDENCE_SCOPE_MISMATCH")
    if any(len(set(slide.evidence_refs)) != len(slide.evidence_refs) or not set(slide.evidence_refs) <= evidence for slide in version.slides):
        raise ValueError("EVIDENCE_REFERENCE_MISMATCH")
    return version


def _shape_roles(root):
    result = {}
    for shape in root.findall("p:cSld/p:spTree/p:sp", NS):
        properties = shape.find("p:nvSpPr/p:cNvPr", NS)
        if properties is None:
            raise ValueError("INVALID_NATIVE_SHAPE")
        name = properties.attrib.get("name", "")
        if name.startswith("gezhi:"):
            if name in result:
                raise ValueError("DUPLICATE_NATIVE_ROLE")
            result[name] = shape
        elif any(_paragraphs(shape, "a")):
            raise ValueError("UNEXPECTED_RENDERED_CONTENT")
    return result


def _pptx_projection(parts, roots, version):
    presentation = roots[MAIN["pptx"]]
    size = presentation.find("p:sldSz", NS)
    if size is None or (int(size.attrib["cx"]), int(size.attrib["cy"])) != (WIDTH * EMU_PER_POINT, HEIGHT * EMU_PER_POINT):
        raise ValueError("INVALID_SLIDE_GEOMETRY")
    graph = _graph("pptx", parts, roots)
    local = graph["relationships"].get(MAIN["pptx"], {})
    identifiers = presentation.findall("p:sldIdLst/p:sldId", NS)
    if len(identifiers) != len(version.slides) or len({element.attrib["id"] for element in identifiers}) != len(identifiers):
        raise ValueError("SLIDE_COUNT_MISMATCH")
    ordered = []
    for element in identifiers:
        edge, target = local[element.attrib["{" + NS["r"] + "}id"]]
        if edge.attrib["Type"] != REL + "slide" or target in ordered:
            raise ValueError("INVALID_SLIDE_ORDER")
        ordered.append(target)
    if set(ordered) != {name for name in roots if re.fullmatch(r"ppt/slides/slide(?:[1-9]|1[0-2])\.xml", name)}:
        raise ValueError("UNEXPECTED_SLIDE_PART")
    for page, (name, expected) in enumerate(zip(ordered, version.slides), 1):
        root = roots[name]
        roles = _shape_roles(root)
        boxes = text_boxes(expected.layout)
        if set(roles) != set(boxes):
            raise ValueError("NATIVE_ROLE_MISMATCH")
        projection = {
            "gezhi:title": (expected.title,), "gezhi:source": (expected.source_note,),
            "gezhi:font": (PPTX_FONT_DISCLOSURE,),
        }
        if expected.layout == "two_column":
            projection.update({"gezhi:left": expected.columns[0], "gezhi:right": expected.columns[1]})
        else:
            projection["gezhi:body"] = expected.body
        for role, texts in projection.items():
            shape = roles[role]
            if _paragraphs(shape, "a") != list(texts or ("",)):
                raise ValueError("SLIDE_TEXT_MISMATCH")
            transform = shape.find("p:spPr/a:xfrm", NS)
            offset, extent = transform.find("a:off", NS), transform.find("a:ext", NS)
            geometry = tuple(int(element.attrib[key]) for element, key in ((offset, "x"), (offset, "y"), (extent, "cx"), (extent, "cy")))
            if geometry != tuple(value * EMU_PER_POINT for value in boxes[role]):
                raise ValueError("INVALID_TEXTBOX_GEOMETRY")
            body = shape.find("p:txBody/a:bodyPr", NS)
            if body is None or body.find("a:noAutofit", NS) is None or body.find("a:normAutofit", NS) is not None or body.find("a:spAutofit", NS) is not None:
                raise ValueError("UNSAFE_TEXT_AUTOFIT")
            if role in {"gezhi:body", "gezhi:left", "gezhi:right"}:
                if any(int(body.attrib.get(key, "0")) != 8 * EMU_PER_POINT for key in ("lIns", "rIns", "tIns", "bIns")):
                    raise ValueError("INVALID_TEXTBOX_MARGIN")
                for paragraph in shape.findall("p:txBody/a:p", NS):
                    properties = paragraph.findall("a:pPr/a:defRPr", NS) + paragraph.findall("a:r/a:rPr", NS)
                    if not properties or any(int(prop.attrib.get("sz", "0")) != BODY_FONT_SIZE * 100 for prop in properties):
                        raise ValueError("INVALID_BODY_FONT_SIZE")
                    if any(prop.find("a:ea", NS) is None or prop.find("a:ea", NS).attrib.get("typeface") != PPTX_FONT for prop in properties):
                        raise ValueError("INVALID_EAST_ASIAN_FONT")
                check_density(tuple(texts), boxes[role], BODY_FONT_SIZE, page)
        description = roles["gezhi:title"].find("p:nvSpPr/p:cNvPr", NS).attrib.get("descr", "")
        if json.loads(description) != {"layout": expected.layout, "evidence_refs": [str(item) for item in expected.evidence_refs]}:
            raise ValueError("SLIDE_EVIDENCE_MISMATCH")
        for shape in root.findall("p:cSld/p:spTree/p:sp", NS):
            transform = shape.find("p:spPr/a:xfrm", NS)
            if transform is None:
                raise ValueError("INVALID_NATIVE_SHAPE")
            offset, extent = transform.find("a:off", NS), transform.find("a:ext", NS)
            x, y, cx, cy = tuple(int(element.attrib[key]) for element, key in ((offset, "x"), (offset, "y"), (extent, "cx"), (extent, "cy")))
            if x < 0 or y < 0 or cx <= 0 or cy <= 0 or x + cx > WIDTH * EMU_PER_POINT or y + cy > HEIGHT * EMU_PER_POINT:
                raise ValueError("OFF_CANVAS_SHAPE")
        note_edges = [(edge, target) for edge, target in graph["relationships"].get(name, {}).values() if edge.attrib["Type"] == REL + "notesSlide"]
        if len(note_edges) != 1:
            raise ValueError("MISSING_NOTES")
        note_root = roots[note_edges[0][1]]
        note_bodies = []
        for shape in note_root.findall("p:cSld/p:spTree/p:sp", NS):
            placeholder = shape.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if placeholder is not None and placeholder.attrib.get("type") == "body":
                note_bodies.append(shape)
        if len(note_bodies) != 1 or "\n".join(_paragraphs(note_bodies[0], "a")) != expected.notes:
            raise ValueError("NOTES_TEXT_MISMATCH")
    return ordered


def _docx_projection(roots, version):
    root = roots[MAIN["docx"]]
    if _paragraphs(root, "w") != _docx_expected_paragraphs(version.lesson):
        raise ValueError("LESSON_TEXT_MISMATCH")
    fonts = root.findall(".//w:rPr/w:rFonts", NS)
    names = {font.attrib.get("{" + NS["w"] + "}eastAsia") for font in fonts}
    if not fonts or not {"黑体", "宋体"} <= names:
        raise ValueError("INVALID_EAST_ASIAN_FONT")
    normal = roots["word/styles.xml"].find("w:style[@w:styleId='Normal']/w:rPr/w:rFonts", NS)
    if normal is None or normal.attrib.get("{" + NS["w"] + "}eastAsia") != "宋体":
        raise ValueError("INVALID_EAST_ASIAN_STYLE")
    tables = root.findall("w:body/w:tbl", NS)
    if len(tables) != 2:
        raise ValueError("LESSON_TABLE_MISMATCH")
    flow_rows = tables[1].findall("w:tr", NS)
    if len(flow_rows) != len(version.lesson.teaching_flow) + 1:
        raise ValueError("LESSON_FLOW_MISMATCH")
    for row, stage in zip(flow_rows[1:], version.lesson.teaching_flow):
        values = ["\n".join(_paragraphs(cell, "w")) for cell in row.findall("w:tc", NS)]
        if values != [stage.stage, str(stage.minutes), stage.content]:
            raise ValueError("LESSON_FLOW_MISMATCH")


def _reopen(kind, raw, version):
    """Supplementary trusted library view; no unsafe candidate reaches this call."""
    if kind == "pptx":
        from pptx import Presentation
        document = Presentation(BytesIO(raw))
        if len(document.slides) != len(version.slides):
            raise ValueError("LIBRARY_SLIDE_MISMATCH")
        for slide, expected in zip(document.slides, version.slides):
            if not slide.has_notes_slide or slide.notes_slide.notes_text_frame is None:
                raise ValueError("LIBRARY_NOTES_MISMATCH")
            if slide.notes_slide.notes_text_frame.text.replace("\v", "\n") != expected.notes:
                raise ValueError("LIBRARY_NOTES_MISMATCH")
            titles = [shape.text_frame.paragraphs[0].text for shape in slide.shapes if shape.name == "gezhi:title"]
            if titles != [expected.title]:
                raise ValueError("LIBRARY_TEXT_MISMATCH")
    else:
        from docx import Document
        document = Document(BytesIO(raw))
        texts = []
        for child in document.element.body:
            for paragraph in child.iter("{" + NS["w"] + "}p"):
                texts.append(_paragraph(paragraph, "w"))
        if texts != _docx_expected_paragraphs(version.lesson):
            raise ValueError("LIBRARY_TEXT_MISMATCH")


def validate_office_bytes(kind, raw: bytes, version: PackageVersionDTO) -> ValidationSummary:
    """Read-only verdict. Invalid bytes yield bounded controlled codes, no paths/text."""
    started = monotonic()
    try:
        if kind not in MAIN:
            raise ValueError("UNSUPPORTED_OFFICE_KIND")
        version = _checked_version(version)
        parts, roots = _parts_and_roots(kind, raw)
        if kind == "pptx":
            _pptx_projection(parts, roots, version)
        else:
            _docx_projection(roots, version)
        _check_stage_time(started)
        _reopen(kind, raw, version)
        _check_stage_time(started)
    except (ValueError, TypeError, KeyError, AttributeError, BadZipFile, ET.ParseError, UnicodeError, OverflowError, NotImplementedError, EOFError, DeflateError) as error:
        code = str(error) if type(error) is ValueError and re.fullmatch(r"[A-Z][A-Z0-9_]{0,63}", str(error)) else "INVALID_OFFICE_PACKAGE"
        return ValidationSummary(valid=False, checks=(code,), warnings=(WARNING,))
    return ValidationSummary(valid=True, checks=("BOUNDED_ZIP_XML", "OWNER_LOCAL_GRAPH", "FROZEN_DIGEST_EVIDENCE", "EXACT_NATIVE_CONTENT", "SAFE_LIBRARY_REOPEN"), warnings=(WARNING,))
